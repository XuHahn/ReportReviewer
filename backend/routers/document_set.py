"""Document Set API router."""

import asyncio
import hashlib
import json as json_mod
import io
import os
import re
import stat
import subprocess
import time
import uuid
import zipfile
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

from fastapi import APIRouter, Body, Depends, HTTPException, Query, Request, Response, UploadFile, File, Form
from fastapi.responses import StreamingResponse

from auth import get_current_user, require_role
from database import save_audit_log_async
from models import (
    User, DocType, DOC_TYPES, SetDocumentCreate,
    DocumentSetStandardsUpdate, DocumentSetProjectGroupUpdate,
)
from services.document_set import DocumentSetService
from utils.logger import get_logger
from utils.pdf_utils import extract_text
from utils.upload_utils import read_upload_limited
from config import (
    EXTRACTION_TOTAL_TIMEOUT_SEC,
    MAX_CONCURRENT_EXTRACTIONS,
    MAX_ZIP_FILES,
    MAX_ZIP_UNCOMPRESSED_BYTES,
    MAX_ZIP_COMPRESSION_RATIO,
    get_max_upload_bytes,
)
from rate_limit import check_upload_async as _rate_check
import database

logger = get_logger("document_set")

router = APIRouter(prefix="/sets", tags=["document_sets"])

# ── Global extraction concurrency control ───────────────────────────────
_extraction_semaphore = asyncio.Semaphore(MAX_CONCURRENT_EXTRACTIONS)
_running_extractions: dict[str, asyncio.Task] = {}


def _evidence_anchor_signature(item) -> str:
    """Hash every visual input while treating lazy enrichment as cache-neutral."""
    metadata = dict(item.metadata)
    source_anchor = metadata.get("source_anchor")
    preview_enriched = bool(
        isinstance(source_anchor, dict)
        and str(source_anchor.get("method") or "").startswith("preview_")
    )
    if preview_enriched:
        metadata.pop("source_anchor", None)
    payload = {
        "quote": item.exact_quote,
        "bbox": [] if preview_enriched else item.bbox,
        "metadata": metadata,
        "extraction_method": getattr(item, "extraction_method", ""),
    }
    return hashlib.sha256(json_mod.dumps(
        payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"),
    ).encode()).hexdigest()


async def _run_unified_review_background(
    *, set_id: str, graph_id: str, gateway, store,
) -> None:
    """Keep document-set lifecycle aligned with the evidence-graph run."""
    from services.unified_review_pipeline import run_unified_review

    try:
        await run_unified_review(
            set_id=set_id, graph_id=graph_id, gateway=gateway,
            store=store, precreated=True,
        )
    except asyncio.CancelledError:
        await database.transition_set_status_async(set_id, "locked", ("reviewing",))
        raise
    except Exception:
        # The failed graph run retains its terminal diagnostics. Keep the
        # document set retryable instead of stranding it in a dead-end state.
        await database.transition_set_status_async(set_id, "locked", ("reviewing",))
        raise
    else:
        await database.transition_set_status_async(set_id, "reviewed", ("reviewing",))


@router.post("/{set_id}/review-runs", status_code=202)
async def start_unified_review_run(
    set_id: str,
    request: Request = None,
    user: User = Depends(require_role("admin", "reviewer")),
):
    """Start the single evidence-graph review engine."""
    await _check_ownership(set_id, user)
    document_set = await database.get_document_set_async(set_id)
    if not document_set:
        raise HTTPException(404, detail=f"文档集不存在: {set_id}")
    latest: dict[str, dict] = {}
    for document in document_set.get("documents", []):
        doc_type = document.get("doc_type", "")
        current = latest.get(doc_type)
        if current is None or int(document.get("doc_version") or 0) > int(current.get("doc_version") or 0):
            latest[doc_type] = document
    required = {"order_form", "test_plan", "original_records", "final_report"}
    missing = sorted(required - set(latest))
    if missing:
        raise HTTPException(400, detail=f"缺少审核所需文档: {', '.join(missing)}")

    from services.evidence_graph_review_engine import EvidenceGraphReviewEngine
    from services.evidence_graph_store import EvidenceGraphStore
    from services.unified_model_gateway import UnifiedModelGateway, UnifiedModelUnavailable
    from utils.background_tasks import start_background_task

    gateway = UnifiedModelGateway()
    unavailable: list[str] = []
    for task in ("structured_extraction", "visual_extraction"):
        try:
            gateway.resolve(task)
        except UnifiedModelUnavailable as exc:
            unavailable.append(str(exc))
    if unavailable:
        raise HTTPException(503, detail={
            "message": "审核所需模型服务尚未配置完整",
            "unavailable": unavailable,
        })
    preflight_errors = await gateway.preflight(("structured_extraction", "visual_extraction"))
    if preflight_errors:
        raise HTTPException(503, detail={
            "message": "审核所需模型服务未就绪",
            "unavailable": preflight_errors,
        })

    store = EvidenceGraphStore(database.DB_PATH)
    store.init_schema()
    active = [
        item for item in store.list_runs(set_id)
        if item["status"] in {"created", "building", "running"}
    ]
    if active:
        raise HTTPException(409, detail="该文档集已有审核任务正在运行")

    original_status = str(document_set.get("status") or "")
    claimed = await database.transition_set_status_async(
        set_id, "reviewing", ("locked", "reviewed"),
    )
    if not claimed:
        raise HTTPException(409, detail="请先完成资料核查并锁定文档集")

    graph_id = f"graph-{uuid.uuid4().hex[:20]}"
    selected_standards = await database.get_document_set_standards_async(set_id)
    try:
        EvidenceGraphReviewEngine(store).create_run(
            graph_id,
            set_id,
            standard_release_ids=[
                standard.selected_release_id for standard in selected_standards
                if standard.selected_release_id
            ],
            model_manifest={
                **gateway.manifest(),
                "extraction_strategy": "native_plus_targeted_qwen",
            },
        )
    except Exception:
        await database.transition_set_status_async(
            set_id, original_status, ("reviewing",),
        )
        raise
    start_background_task(
        _run_unified_review_background(
            set_id=set_id, graph_id=graph_id, gateway=gateway, store=store,
        ),
        name=f"review-{graph_id}",
        key=f"unified-review:{set_id}",
    )
    await save_audit_log_async(
        client_ip=_client_ip(request),
        action="unified_review_started",
        filename=set_id,
        employee_id=user.employee_id,
        detail=json_mod.dumps({"set_id": set_id, "graph_id": graph_id}),
    )
    return {"set_id": set_id, "run_id": graph_id, "status": "building"}


@router.get("/{set_id}/review-runs")
async def list_unified_review_runs(
    set_id: str,
    user: User = Depends(require_role("admin", "reviewer", "viewer")),
):
    await _check_ownership(set_id, user)
    from services.evidence_graph_store import EvidenceGraphStore

    store = EvidenceGraphStore(database.DB_PATH)
    store.init_schema()
    runs = await asyncio.to_thread(store.list_runs, set_id)
    return {"runs": runs, "total": len(runs)}


@router.get("/{set_id}/review-runs/{graph_id}")
async def get_unified_review_run(
    set_id: str,
    graph_id: str,
    user: User = Depends(require_role("admin", "reviewer", "viewer")),
):
    await _check_ownership(set_id, user)
    from services.evidence_graph_store import EvidenceGraphStore

    store = EvidenceGraphStore(database.DB_PATH)
    try:
        snapshot = await asyncio.to_thread(store.get_snapshot, graph_id)
    except KeyError:
        raise HTTPException(404, detail="审核运行不存在")
    if snapshot.run.set_id != set_id:
        raise HTTPException(404, detail="审核运行不存在")
    return snapshot.model_dump(mode="json")


@router.get("/{set_id}/review-runs/{graph_id}/evidence/{evidence_id}/preview")
async def preview_unified_review_evidence(
    set_id: str,
    graph_id: str,
    evidence_id: str,
    request: Request,
    related_evidence_ids: str = Query("", max_length=4096),
    highlight_kinds: str = Query("", max_length=512),
    focus_index: int | None = Query(None, ge=0, le=11),
    user: User = Depends(require_role("admin", "reviewer", "viewer")),
):
    """Render one source page, optionally combining same-page evidence anchors."""
    await _check_ownership(set_id, user)
    import fitz
    from services.document_preview import (
        evidence_preview_cache_key,
        extract_archive_member,
        load_evidence_preview_cache,
        render_evidence_page_cached,
    )
    from services.evidence_graph_store import EvidenceGraphStore

    store = EvidenceGraphStore(database.DB_PATH)
    runs = await asyncio.to_thread(store.list_runs, set_id)
    if not any(run["graph_id"] == graph_id for run in runs):
        raise HTTPException(404, detail="审核运行不存在")
    evidence = await asyncio.to_thread(store.get_evidence, graph_id, evidence_id)
    if evidence is None:
        raise HTTPException(404, detail="证据不存在")
    if evidence.page_number < 1:
        raise HTTPException(422, detail="该证据没有可预览的页面定位")

    requested_ids = list(dict.fromkeys([
        item.strip() for item in related_evidence_ids.split(",")
        if item.strip() and item.strip() != evidence_id
    ]))[:11]
    related = []
    primary_member = re.search(r"member-(\d+)", str(evidence.metadata.get("unit_id") or ""))
    primary_source_key = (
        evidence.doc_id, evidence.doc_type, evidence.filename, evidence.page_number,
        primary_member.group(1) if primary_member else "",
        str(evidence.metadata.get("release_id") or ""),
    )
    for requested_id in requested_ids:
        item = await asyncio.to_thread(store.get_evidence, graph_id, requested_id)
        if item is None:
            continue
        item_member = re.search(r"member-(\d+)", str(item.metadata.get("unit_id") or ""))
        item_source_key = (
            item.doc_id, item.doc_type, item.filename, item.page_number,
            item_member.group(1) if item_member else "",
            str(item.metadata.get("release_id") or ""),
        )
        if item_source_key == primary_source_key:
            related.append(item)
    page_evidence = [evidence, *related]
    kinds = [item.strip() for item in highlight_kinds.split(",")]
    allowed_kinds = {"error", "warning", "success", "reference"}
    highlights = []
    for index, item in enumerate(page_evidence):
        preview_metadata = {**item.metadata, "_extraction_method": item.extraction_method}
        if item.extraction_method == "deterministic_missing_field_region":
            preview_metadata.setdefault("role", "missing_field")
        highlights.append({
            "quote": item.exact_quote,
            "bbox": item.bbox,
            "metadata": preview_metadata,
            "kind": (
                kinds[index]
                if index < len(kinds) and kinds[index] in allowed_kinds else "error"
            ),
        })

    cache_key = evidence_preview_cache_key({
        "graph_id": graph_id,
        "source_key": primary_source_key,
        "page_number": evidence.page_number,
        "evidence_ids": [item.evidence_id for item in page_evidence],
        "content_hashes": [item.content_hash for item in page_evidence],
        # Locator enrichment and corrected quotes must invalidate an existing
        # render even when the immutable evidence id stays the same.
        "anchor_signatures": [_evidence_anchor_signature(item) for item in page_evidence],
        "highlight_kinds": [item["kind"] for item in highlights],
        "highlight_labels": [
            str(item["metadata"].get("role") or item["metadata"].get("verdict") or "")
            for item in highlights
        ],
        "focus_index": focus_index,
        "render_scale": os.getenv("UNITIZER_RENDER_SCALE", "2.0"),
    })
    etag = f'"{cache_key}"'
    if request.headers.get("if-none-match", "").strip() == etag:
        logger.info(
            "evidence_page_preview_not_modified",
            graph_id=graph_id,
            evidence_id=evidence_id,
        )
        return Response(
            status_code=304,
            headers={
                "Cache-Control": "private, no-cache, max-age=0, must-revalidate",
                "ETag": etag,
                "X-Content-Type-Options": "nosniff",
            },
        )

    preview = await asyncio.to_thread(load_evidence_preview_cache, cache_key)
    cache_hit = preview is not None

    if preview is None:
        filename = evidence.filename
        file_bytes: bytes | None = None
        if evidence.doc_type == "test_standard":
            release_id = str(evidence.metadata.get("release_id") or "")
            release = await database.get_standard_release_async(release_id) if release_id else None
            standard_file = await database.get_standard_file_async(str((release or {}).get("standard_id") or ""))
            if standard_file:
                filename, file_bytes = standard_file
        else:
            await _require_document_in_set(set_id, evidence.doc_id)
            file_bytes = await database.get_file_content_async(evidence.doc_id)
            member_match = re.search(r"member-(\d+)", str(evidence.metadata.get("unit_id") or ""))
            if file_bytes is not None and member_match:
                try:
                    filename, file_bytes = await asyncio.to_thread(
                        extract_archive_member, file_bytes, int(member_match.group(1)),
                    )
                except (IndexError, zipfile.BadZipFile, ValueError, OSError):
                    raise HTTPException(422, detail="无法读取该证据对应的原始记录文件")
        if file_bytes is None:
            raise HTTPException(404, detail="证据原文件不存在或已被清理")

        try:
            preview = await asyncio.to_thread(
                render_evidence_page_cached,
                cache_key,
                file_bytes,
                filename,
                evidence.page_number,
                quote=evidence.exact_quote,
                bbox=evidence.bbox,
                metadata=evidence.metadata,
                highlights=highlights,
                focus_index=focus_index,
            )
            cache_hit = preview.cache_hit
        except IndexError:
            raise HTTPException(422, detail="证据页码超出原文件范围")
        except (ValueError, OSError, RuntimeError, fitz.FileDataError, subprocess.SubprocessError) as exc:
            logger.warning(
                "evidence_page_preview_failed",
                graph_id=graph_id,
                evidence_id=evidence_id,
                doc_type=evidence.doc_type,
                page_number=evidence.page_number,
                error_type=type(exc).__name__,
            )
            raise HTTPException(503, detail="当前证据页暂时无法生成视觉预览")

    enriched_count = 0
    for index, item in enumerate(page_evidence):
        resolved_bbox = (
            preview.resolved_bboxes[index]
            if index < len(preview.resolved_bboxes) else None
        )
        strategy = (
            preview.locator_strategies[index]
            if index < len(preview.locator_strategies) else "none"
        )
        resolved_rectangles = (
            preview.resolved_rectangles[index]
            if index < len(preview.resolved_rectangles) else ()
        )
        if item.bbox or resolved_bbox is None or strategy not in {
            "exact_quote", "normalized_quote", "structured_comparison_row",
            "structured_field_row", "structured_parameter_value",
            "structured_missing_field", "semantic_structured_row",
            "semantic_parameter_ordinal_row", "strict_ordered_quote",
            "structured_metadata_row",
        }:
            continue
        enriched = await asyncio.to_thread(
            store.enrich_evidence_locator,
            graph_id,
            item.evidence_id,
            list(resolved_bbox),
            method=f"preview_{strategy}",
            rectangles=[list(rectangle) for rectangle in resolved_rectangles],
            rendered_pdf_hash=preview.rendered_pdf_hash,
            page_width=preview.page_width,
            page_height=preview.page_height,
        )
        enriched_count += int(enriched)
    logger.info(
        "evidence_page_preview_rendered",
        graph_id=graph_id,
        evidence_id=evidence_id,
        doc_type=evidence.doc_type,
        page_number=evidence.page_number,
        highlight_strategy=preview.highlight_strategy,
        highlight_count=preview.highlight_count,
        evidence_count=len(page_evidence),
        focus_index=focus_index,
        focus_applied=preview.focus_applied,
        cache_hit=cache_hit,
        locator_enriched_count=enriched_count,
    )
    return Response(
        content=preview.content,
        media_type="image/png",
        headers={
            "Content-Disposition": "inline",
            # Private clients may store the immutable graph representation but
            # must revalidate; the ETag includes the renderer contract version.
            "Cache-Control": "private, no-cache, max-age=0, must-revalidate",
            "ETag": etag,
            "X-Preview-Page-Count": str(preview.page_count),
            "X-Evidence-Highlight-Count": str(preview.highlight_count),
            "X-Evidence-Highlight-Strategy": preview.highlight_strategy,
            "X-Evidence-Record-Count": str(len(page_evidence)),
            "X-Evidence-Focus-Applied": "1" if preview.focus_applied else "0",
            "X-Evidence-Preview-Cache": "hit" if cache_hit else "miss",
            "X-Evidence-Locator-Enriched": str(enriched_count),
            "X-Evidence-Anchor-Y": ",".join(
                "" if position is None else f"{position:.6f}"
                for position in preview.anchor_y_positions
            ),
            "X-Content-Type-Options": "nosniff",
        },
    )


@router.get("/{set_id}/review-runs/{graph_id}/export")
async def export_unified_review_run(
    set_id: str,
    graph_id: str,
    request: Request,
    format: str = Query("xlsx", pattern="^(xlsx|pdf|evidence)$"),
    user: User = Depends(require_role("admin", "reviewer", "viewer")),
):
    await _check_ownership(set_id, user)
    from services.evidence_graph_exporter import (
        generate_evidence_package,
        generate_review_excel,
        generate_review_pdf,
    )
    from services.evidence_graph_store import EvidenceGraphStore

    try:
        snapshot = await asyncio.to_thread(
            EvidenceGraphStore(database.DB_PATH).get_snapshot, graph_id,
        )
    except KeyError:
        raise HTTPException(404, detail="审核运行不存在")
    if snapshot.run.set_id != set_id:
        raise HTTPException(404, detail="审核运行不存在")

    generators = {
        "xlsx": (generate_review_excel, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
        "pdf": (generate_review_pdf, "application/pdf", "pdf"),
        "evidence": (generate_evidence_package, "application/zip", "zip"),
    }
    generator, media_type, suffix = generators[format]
    try:
        content = await asyncio.to_thread(generator, snapshot)
    except (OSError, subprocess.SubprocessError) as exc:
        logger.error(
            "unified_review_export_failed",
            graph_id=graph_id,
            export_format=format,
            error_type=type(exc).__name__,
        )
        raise HTTPException(503, detail="导出服务暂不可用")
    await save_audit_log_async(
        client_ip=_client_ip(request),
        action="unified_review_exported",
        filename=set_id,
        employee_id=user.employee_id,
        detail=json_mod.dumps({"set_id": set_id, "graph_id": graph_id, "format": format}),
    )
    filename = f"{set_id}-{graph_id}-review.{suffix}"
    return StreamingResponse(
        io.BytesIO(content),
        media_type=media_type,
        headers={"Content-Disposition": _encoded_attachment_filename(filename)},
    )


@router.put("/{set_id}/review-runs/{graph_id}/findings/{finding_id}/decision")
async def decide_unified_review_finding(
    set_id: str,
    graph_id: str,
    finding_id: str,
    request: Request,
    body: dict = Body(...),
    user: User = Depends(require_role("admin", "reviewer")),
):
    await _check_ownership(set_id, user)
    from services.evidence_graph_models import FindingDecision
    from services.evidence_graph_store import EvidenceGraphStore

    store = EvidenceGraphStore(database.DB_PATH)
    store.init_schema()
    try:
        snapshot = await asyncio.to_thread(store.get_snapshot, graph_id)
    except KeyError:
        raise HTTPException(404, detail="审核运行不存在")
    if snapshot.run.set_id != set_id:
        raise HTTPException(404, detail="审核运行不存在")
    decision_value = str(body.get("decision") or "")
    resolution_code = str(body.get("resolution_code") or "")
    resolution_status = str(body.get("resolution_status") or "")
    comment = str(body.get("comment") or "").strip()
    resolution_contract = {
        "report_revision": ("confirmed", "open"),
        "raw_record_supplement": ("confirmed", "open"),
        "source_correction": ("confirmed", "open"),
        "not_applicable": ("dismissed", "not_applicable"),
        "false_positive": ("dismissed", "dismissed"),
        "deferred": ("unresolved", "deferred"),
        "other": ("advisory", "open"),
    }
    expected = resolution_contract.get(resolution_code)
    if expected and (decision_value, resolution_status) != expected:
        raise HTTPException(422, detail="处理方式与人工结论不一致")
    if (decision_value in {"confirmed", "dismissed", "advisory"} or resolution_code) and not comment:
        raise HTTPException(422, detail="人工裁决必须填写理由")
    try:
        decision = await asyncio.to_thread(store.set_finding_decision, FindingDecision(
            graph_id=graph_id, finding_id=finding_id, decision=decision_value,
            resolution_code=resolution_code, resolution_status=resolution_status,
            comment=comment, actor_id=user.employee_id,
        ))
    except ValueError as exc:
        raise HTTPException(422, detail=str(exc))
    store.record_event(
        graph_id, "finding_human_decision", actor_type="human",
        actor_id=user.employee_id, subject_type="review_finding",
        subject_id=finding_id, new_state=decision_value,
        reason_code=resolution_code or "human_review",
        payload={
            "comment_length": len(comment),
            "resolution_code": resolution_code,
            "resolution_status": resolution_status,
        },
    )
    await save_audit_log_async(
        client_ip=_client_ip(request), action="unified_finding_decided",
        filename=set_id, employee_id=user.employee_id,
        detail=json_mod.dumps({"graph_id": graph_id, "finding_id": finding_id,
                               "decision": decision_value,
                               "resolution_code": resolution_code,
                               "resolution_status": resolution_status,
                               "comment_length": len(comment)}),
    )
    return decision.model_dump(mode="json")


async def shutdown_extraction_tasks() -> None:
    tasks = list(_running_extractions.values())
    for task in tasks:
        task.cancel()
    if tasks:
        await asyncio.gather(*tasks, return_exceptions=True)
    _running_extractions.clear()


def _client_ip(request: Request | None) -> str:
    return request.client.host if request and request.client else "unknown"


def _encoded_attachment_filename(filename: str) -> str:
    encoded = quote(filename, safe="")
    return f"attachment; filename=\"{encoded}\"; filename*=UTF-8''{encoded}"


def _source_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, dict):
        parts = []
        for key, val in value.items():
            if isinstance(val, (dict, list)):
                rendered = json_mod.dumps(val, ensure_ascii=False)
            else:
                rendered = str(val)
            parts.append(f"{key}: {rendered}")
        return "\n".join(parts)
    if isinstance(value, list):
        return "\n".join(str(v) for v in value)
    return str(value)


def _generate_document_set_audit_excel(
    set_id: str,
    overview: dict | None,
    issues: list[dict],
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    severity_label = {
        "CRITICAL": "严重",
        "WARNING": "警告",
        "INFO": "提示",
    }
    human_status_label = {
        "pending": "待处理",
        "confirmed": "已确认",
        "ignored": "已忽略",
        "false_positive": "非错误",
        "needs_review": "需复核",
    }
    severity_fill = {
        "CRITICAL": "FEE2E2",
        "WARNING": "FEF3C7",
        "INFO": "DBEAFE",
    }
    severity_font = {
        "CRITICAL": "DC2626",
        "WARNING": "D97706",
        "INFO": "2563EB",
    }
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "问题明细"
    headers = [
        "序号", "级别", "问题类别", "检查编号", "问题字段",
        "问题描述", "人工状态", "人工备注", "标注人", "标注时间",
        "来源", "来源证据", "创建时间",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, issue in enumerate(issues, 1):
        sev = str(issue.get("severity", ""))
        status = str(issue.get("human_status", "pending"))
        row = [
            idx,
            severity_label.get(sev, sev),
            issue.get("category", ""),
            issue.get("check_id", ""),
            issue.get("field_name", ""),
            issue.get("description", ""),
            human_status_label.get(status, status),
            issue.get("human_comment", ""),
            issue.get("annotated_by", ""),
            issue.get("annotated_at", ""),
            issue.get("source", ""),
            _source_text(issue.get("sources", {})),
            issue.get("created_at", ""),
        ]
        ws.append(row)
        excel_row = idx + 1
        for cell in ws[excel_row]:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(excel_row, 2).fill = PatternFill(
            start_color=severity_fill.get(sev, "FFFFFF"),
            end_color=severity_fill.get(sev, "FFFFFF"),
            fill_type="solid",
        )
        ws.cell(excel_row, 2).font = Font(
            color=severity_font.get(sev, "334155"),
            bold=True,
        )

    widths = [8, 10, 16, 12, 24, 60, 12, 28, 14, 20, 14, 72, 20]
    for col_idx, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, col_idx).column_letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    summary = wb.create_sheet("统计汇总")
    summary_rows = [
        ["文档集编号", set_id],
        ["导出时间", datetime.now().isoformat(timespec="seconds")],
        ["问题总数", len(issues)],
        ["严重", sum(1 for i in issues if i.get("severity") == "CRITICAL")],
        ["警告", sum(1 for i in issues if i.get("severity") == "WARNING")],
        ["提示", sum(1 for i in issues if i.get("severity") == "INFO")],
        ["待处理", sum(1 for i in issues if i.get("human_status") == "pending")],
        ["已确认", sum(1 for i in issues if i.get("human_status") == "confirmed")],
        ["已忽略", sum(1 for i in issues if i.get("human_status") == "ignored")],
    ]
    for row in summary_rows:
        summary.append(row)
    for row in summary.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        row[0].font = Font(bold=True)
    summary.column_dimensions["A"].width = 18
    summary.column_dimensions["B"].width = 42

    files_sheet = wb.create_sheet("文件清单")
    files_sheet.append(["文件类型", "文件名", "版本", "大小(KB)", "提取状态", "上传时间"])
    for cell in files_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for doc in (overview or {}).get("files", []) or []:
        files_sheet.append([
            doc.get("doc_type", ""),
            doc.get("filename", ""),
            doc.get("doc_version", ""),
            doc.get("file_size_kb", ""),
            doc.get("extraction_status", ""),
            doc.get("created_at", ""),
        ])
    for row in files_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col, width in zip(["A", "B", "C", "D", "E", "F"], [18, 44, 10, 12, 14, 22]):
        files_sheet.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _check_ownership(set_id: str, user: User):
    """Raise 403 if user is not admin and does not own the set."""
    if user.role == "admin":
        return
    ds = await database.get_document_set_async(set_id)
    if not ds:
        raise HTTPException(404, detail=f"文档集不存在: {set_id}")
    if ds.get("employee_id", "") != user.employee_id:
        raise HTTPException(403, detail="无权操作此文档集")


async def _require_document_in_set(set_id: str, doc_id: str) -> None:
    if not await database.document_belongs_to_set_async(set_id, doc_id):
        raise HTTPException(404, detail=f"文档不存在: {doc_id}")


async def _require_editable_set(set_id: str) -> dict:
    """Require the document set to be an unlocked draft before mutating sources."""
    document_set = await database.get_document_set_async(set_id)
    if not document_set:
        raise HTTPException(404, detail=f"文档集不存在: {set_id}")
    if document_set.get("status") not in {"incomplete", "revision"}:
        raise HTTPException(409, detail="文档集已锁定；请先创建修订版本再修改资料")
    return document_set




@router.post("")
async def create_set(request: Request = None, user: User = Depends(require_role("admin", "reviewer"))):
    try:
        ds = await DocumentSetService.create_set(user.employee_id)
        logger.info("create_set", set_id=ds.set_id, user=user.employee_id,
                     extra={"status": ds.status})
        await save_audit_log_async(
            client_ip=_client_ip(request), action="set_create",
            filename=ds.set_id, employee_id=user.employee_id,
            detail=json_mod.dumps({"status": ds.status}))
        return {"set_id": ds.set_id, "status": ds.status, "employee_id": ds.employee_id}
    except Exception as e:
        logger.error("create_set_failed", user=user.employee_id, error=str(e))
        raise


@router.get("")
async def list_sets(user: User = Depends(require_role("admin", "reviewer", "viewer"))):
    emp = "" if user.role == "admin" else user.employee_id
    sets = await DocumentSetService.list_sets(emp)
    logger.info("list_sets", user=user.employee_id, extra={"count": len(sets), "role": user.role})
    return {"sets": sets, "total": len(sets)}


@router.get("/stats/overview")
async def get_stats_overview(user: User = Depends(require_role("admin", "reviewer", "viewer"))):
    """Get aggregate statistics from document sets and evidence-graph findings."""
    try:
        stats = await database.get_set_stats_async(
            "" if user.role == "admin" else user.employee_id,
        )
        logger.info("set_stats", user=user.employee_id,
                    extra={"total_sets": stats["total_sets"]})
        return stats
    except Exception as e:
        logger.error("set_stats_failed", user=user.employee_id, error=str(e))
        raise


@router.get("/{set_id}")
async def get_set(set_id: str, user: User = Depends(require_role("admin", "reviewer", "viewer"))):
    await _check_ownership(set_id, user)
    overview = await DocumentSetService.get_set_overview(set_id)
    logger.info("get_set", set_id=set_id, user=user.employee_id,
                 extra={"found": overview is not None})
    if not overview:
        raise HTTPException(404, detail=f"文档集不存在: {set_id}")
    return overview


@router.get("/{set_id}/standards")
async def get_set_standards(
    set_id: str,
    user: User = Depends(require_role("admin", "reviewer", "viewer")),
):
    await _check_ownership(set_id, user)
    standards = await database.get_document_set_standards_async(set_id)
    return {
        "set_id": set_id,
        "standards": [standard.model_dump() for standard in standards],
        "standard_review_enabled": bool(standards),
    }


@router.put("/{set_id}/project-group")
async def update_set_project_group(
    set_id: str,
    body: DocumentSetProjectGroupUpdate,
    request: Request = None,
    user: User = Depends(require_role("admin", "reviewer")),
):
    await _check_ownership(set_id, user)
    if body.project_group_id:
        group = await database.get_project_group_async(body.project_group_id)
        if not group:
            raise HTTPException(422, detail="项目组不存在")
        if user.role != "admin" and user.employee_id not in group.get("members", []):
            raise HTTPException(403, detail="只能绑定自己所属的项目组")
    try:
        result = await database.set_document_set_project_group_async(
            set_id, body.project_group_id,
        )
    except ValueError as exc:
        raise HTTPException(422, detail=str(exc)) from exc
    await save_audit_log_async(
        client_ip=_client_ip(request), action="set_project_group_update",
        filename=set_id, employee_id=user.employee_id,
        detail=json_mod.dumps(result, ensure_ascii=False),
    )
    return {"set_id": set_id, **result}


@router.delete("/{set_id}")
async def delete_set(
    set_id: str,
    request: Request = None,
    user: User = Depends(require_role("admin")),
):
    ds = await database.get_document_set_async(set_id)
    if not ds:
        raise HTTPException(404, detail="文档集不存在")
    cancelled_tasks: list[asyncio.Task] = []
    for task_key, task in list(_running_extractions.items()):
        if task_key.startswith(f"{set_id}:"):
            task.cancel()
            cancelled_tasks.append(task)
            _running_extractions.pop(task_key, None)
    if cancelled_tasks:
        await asyncio.gather(*cancelled_tasks, return_exceptions=True)
    deleted = await database.delete_document_set_async(set_id)
    if not deleted:
        raise HTTPException(404, detail="文档集不存在")
    await save_audit_log_async(
        client_ip=_client_ip(request), action="set_delete", filename=set_id,
        employee_id=user.employee_id,
        detail=json_mod.dumps({"previous_status": ds.get("status", "")}),
    )
    logger.warning("document_set_deleted", set_id=set_id, user=user.employee_id)
    return {"deleted": True, "set_id": set_id}


@router.get("/{set_id}/standard-references")
async def get_set_standard_references(
    set_id: str,
    user: User = Depends(require_role("admin", "reviewer", "viewer")),
):
    await _check_ownership(set_id, user)
    if not await DocumentSetService.get_set_overview(set_id):
        raise HTTPException(404, detail=f"文档集不存在: {set_id}")
    return {"set_id": set_id, **await database.get_document_set_standard_resolution_async(set_id)}


@router.put("/{set_id}/standards")
async def update_set_standards(
    set_id: str,
    body: DocumentSetStandardsUpdate,
    request: Request = None,
    user: User = Depends(require_role("admin", "reviewer")),
):
    await _check_ownership(set_id, user)
    overview = await DocumentSetService.get_set_overview(set_id)
    if not overview:
        raise HTTPException(404, detail=f"文档集不存在: {set_id}")
    if overview.get("status") not in {"incomplete", "revision"}:
        raise HTTPException(400, detail="文档集锁定后不能修改所选标准")
    try:
        standards = await database.set_document_set_standards_async(
            set_id, body.standard_ids, user.employee_id,
        )
        skips = await database.set_document_set_standard_skips_async(
            set_id, body.skips, user.employee_id,
        )
    except ValueError as exc:
        raise HTTPException(422, detail=str(exc)) from exc
    logger.info(
        "set_standard_selection_updated", set_id=set_id,
        standard_ids=body.standard_ids, user=user.employee_id,
    )
    await save_audit_log_async(
        client_ip=_client_ip(request), action="set_standard_selection",
        filename=set_id, employee_id=user.employee_id,
        detail=json_mod.dumps({"standard_ids": body.standard_ids,
                               "skipped_reference_count": len(skips)}, ensure_ascii=False),
    )
    return {
        "set_id": set_id,
        "standards": [standard.model_dump() for standard in standards],
        "standard_skips": skips,
        "standard_review_enabled": bool(standards),
    }


@router.post("/{set_id}/documents")
async def add_document(
    set_id: str,
    file: UploadFile = File(...),
    doc_type: str = Form(...),
    replace_doc_id: str = Form(""),
    request: Request = None,
    user: User = Depends(require_role("admin", "reviewer")),
):
    await _check_ownership(set_id, user)
    await _require_editable_set(set_id)
    if doc_type not in DOC_TYPES:
        raise HTTPException(422, detail=f"无效的文档类型: {doc_type}")
    if replace_doc_id:
        await _require_document_in_set(set_id, replace_doc_id)

    client_ip = _client_ip(request)
    if not await _rate_check(client_ip):
        raise HTTPException(429, detail="上传过于频繁，请稍后再试")
    upload_limit = get_max_upload_bytes()
    try:
        # Read file content (moved inside try/except for proper error handling)
        file_bytes = await read_upload_limited(file, upload_limit)
        file_size_kb = len(file_bytes) // 1024

        # Create document record
        try:
            doc = await DocumentSetService.add_document(
                set_id=set_id, doc_type=doc_type, filename=file.filename or "unknown",
                file_size_kb=file_size_kb, replace_doc_id=replace_doc_id,
            )
        except ValueError as exc:
            raise HTTPException(409, detail=str(exc)) from exc

        # Store original file bytes for later download (打开原始文件)
        await database.store_file_content_async(doc.doc_id, file_bytes)

        logger.info("add_document", set_id=set_id, doc_id=doc.doc_id,
                     user=user.employee_id,
                     extra={"doc_type": doc_type, "filename": file.filename,
                            "file_size_kb": file_size_kb,
                            "replace_doc_id": replace_doc_id})
        await save_audit_log_async(
            client_ip=_client_ip(request), action="document_upload",
            filename=doc.doc_id, file_size_kb=file_size_kb,
            employee_id=user.employee_id,
            detail=json_mod.dumps({"set_id": set_id, "doc_type": doc_type,
                                   "filename": file.filename,
                                   "replace_doc_id": replace_doc_id}))

        # Trigger extraction based on doc_type (fire-and-forget, update status async)
        task = asyncio.create_task(_extract_with_timeout(
            set_id, doc.doc_id, doc_type, file_bytes, file.filename or "",
            user.employee_id, doc.doc_version, 0,
        ))
        _running_extractions[f"{set_id}:{doc.doc_id}"] = task

        return doc.model_dump()
    except HTTPException as exc:
        if exc.status_code == 413:
            logger.warning(
                "document_upload_rejected",
                set_id=set_id,
                user=user.employee_id,
                filename=file.filename,
                limit_bytes=upload_limit,
                reason="file_too_large",
            )
        raise
    except Exception as e:
        logger.error("add_document_failed", set_id=set_id, user=user.employee_id,
                      error=str(e), extra={"doc_type": doc_type, "filename": file.filename})
        raise HTTPException(500, detail=f"文档上传失败: {str(e)}")


@router.delete("/{set_id}/documents/{doc_id}")
async def delete_document(
    set_id: str, doc_id: str,
    request: Request = None,
    user: User = Depends(require_role("admin", "reviewer")),
):
    """Remove a document (and its child versions) from a set."""
    await _check_ownership(set_id, user)
    await _require_editable_set(set_id)
    try:
        ok = await DocumentSetService.delete_document(set_id, doc_id, user.employee_id)
        logger.info("delete_document", set_id=set_id, doc_id=doc_id,
                     user=user.employee_id)
        await save_audit_log_async(
            client_ip=_client_ip(request), action="document_delete",
            filename=doc_id, employee_id=user.employee_id,
            detail=json_mod.dumps({"set_id": set_id}))
        return {"deleted": ok}
    except HTTPException:
        raise
    except Exception as e:
        logger.error("delete_document_failed", set_id=set_id, doc_id=doc_id,
                      user=user.employee_id, error=str(e))
        raise HTTPException(500, detail=f"文档删除失败: {str(e)}")


def _batch_result_to_dict(batch) -> dict:
    """Convert BatchResult dataclass to JSON-serializable dict."""
    from dataclasses import asdict
    return asdict(batch)


def _extract_raw_records_sync(file_bytes: bytes) -> tuple[str, str]:
    """Synchronous raw_records ZIP extraction — runs in thread pool.

    Flow: extract entire ZIP to disk → walk directory for PDFs →
    parse filenames (ASCII structure only) → read each PDF →
    override ALL Chinese metadata from PDF content.

    Design: Chinese text ALWAYS comes from PDF content, never from ZIP
    filenames.  ZIP filenames may use any encoding (GBK, Shift-JIS, UTF-8
    etc.) — we don't attempt to guess.  Instead we rely on the ASCII parts
    of the filename (report_id, test_item_code, sequence, timestamp) for
    structural parsing, then use PDF content as the authoritative source
    for all Chinese fields.
    """
    import zipfile, io, tempfile, os
    from services.document_unitizer import (
        archive_member_manifest,
        display_zip_filename,
    )
    plain_text = ""
    structured_json = ""

    with tempfile.TemporaryDirectory() as tmpdir:
        # Step 1: Extract entire ZIP to disk
        display_names: dict[str, str] = {}
        identities_by_raw_name: dict[str, list[dict]] = {}
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
            members = zf.infolist()
            if len(members) > MAX_ZIP_FILES:
                raise ValueError(f"ZIP 文件数量超过限制: {len(members)}/{MAX_ZIP_FILES}")
            total_uncompressed = sum(max(0, item.file_size) for item in members)
            if total_uncompressed > MAX_ZIP_UNCOMPRESSED_BYTES:
                raise ValueError(
                    f"ZIP 解压后大小超过限制: {total_uncompressed}/{MAX_ZIP_UNCOMPRESSED_BYTES}"
                )
            compressed = sum(max(0, item.compress_size) for item in members)
            ratio = total_uncompressed / max(1, compressed)
            if ratio > MAX_ZIP_COMPRESSION_RATIO:
                raise ValueError(
                    f"ZIP 压缩比异常: {ratio:.1f}/{MAX_ZIP_COMPRESSION_RATIO:.1f}"
                )
            root_path = Path(tmpdir).resolve()
            for item in members:
                target = (root_path / item.filename).resolve()
                if root_path != target and root_path not in target.parents:
                    raise ValueError(f"ZIP 包含不安全路径: {item.filename}")
                if item.flag_bits & 0x1:
                    raise ValueError(f"ZIP 包含加密文件: {item.filename}")
                unix_mode = item.external_attr >> 16
                if unix_mode and stat.S_ISLNK(unix_mode):
                    raise ValueError(f"ZIP 包含符号链接: {item.filename}")
                display_names[item.filename] = display_zip_filename(item)
            for identity in archive_member_manifest(file_bytes):
                identities_by_raw_name.setdefault(
                    str(identity.get("raw_filename") or ""), [],
                ).append(identity)
            zf.extractall(tmpdir)

        # Step 2: Walk disk to find all files
        # decoded display name, full path, stable archive index, member hash
        pdf_entries: list[tuple[str, str, int, str]] = []
        all_names: list[str] = []
        for root, dirs, files in os.walk(tmpdir):
            for f in files:
                if f.startswith('.'):          # skip hidden files
                    continue
                full_path = os.path.join(root, f)
                relative_name = Path(full_path).relative_to(tmpdir).as_posix()
                display_name = (
                    display_names.get(relative_name, relative_name)
                    .replace("\\", "/").rsplit("/", 1)[-1]
                )
                all_names.append(display_name)
                if f.lower().endswith('.pdf'):
                    identities = identities_by_raw_name.get(relative_name, [])
                    identity = identities.pop(0) if identities else {}
                    pdf_entries.append((
                        display_name,
                        full_path,
                        int(identity.get("member_index") or 0),
                        str(identity.get("content_hash") or ""),
                    ))

        if not pdf_entries:
            plain_text = f"ZIP 包含 {len(all_names)} 个文件，未找到 PDF"
            return plain_text, structured_json

        # Sort for deterministic processing order
        pdf_entries.sort(key=lambda item: (item[0], item[2], item[1]))
        pdf_names = [name for name, _path, _index, _hash in pdf_entries]
        all_names.sort()

        plain_text = (
            f"ZIP 包含 {len(all_names)} 个文件 ({len(pdf_names)} 个PDF):\n"
            + "\n".join(all_names[:50])
        )

        from services.raw_records_filename_parser import parse_filename_or_fallback
        from services.raw_records_pdf_extractor import (
            extract_and_populate, override_from_pdf_header,
        )
        from services.raw_records_table_gate import extract_table_inventory
        from services.raw_records_table_extractor import extract_all_tables
        from services.raw_records_instrument_extractor import (
            extract_all_instruments, deduplicate_instruments,
            validate_all_calibrations,
        )
        from services.raw_records_aggregator import (
            BatchResult, summarize_conclusions, validate_extraction,
        )

        # Step 3: Parse filenames — only ASCII segments are reliable.
        # Chinese fields will be overridden from PDF content in step 4.
        metas = [parse_filename_or_fallback(name) for name in pdf_names]
        for meta, (_name, _path, member_index, member_hash) in zip(metas, pdf_entries):
            meta.archive_member_index = member_index
            meta.archive_member_hash = member_hash

        # Step 4: Read each PDF from disk, extract text,
        # override ALL Chinese fields from PDF content
        for m, (_display_name, pdf_path, _member_index, _member_hash) in zip(metas, pdf_entries):
            if not os.path.exists(pdf_path):
                continue
            with open(pdf_path, 'rb') as pf:
                pdf_bytes = pf.read()
            extract_and_populate(m, pdf_bytes)
            m.table_inventory = extract_table_inventory(pdf_bytes, m.filename)
            override_from_pdf_header(m)

        # Step 5: Extract data tables, instruments, conclusions
        all_data_rows = extract_all_tables(metas)
        instruments = extract_all_instruments(metas)
        deduped = deduplicate_instruments(instruments)
        cal_issues = validate_all_calibrations(deduped, metas)
        conclusions = summarize_conclusions(metas)

        validation_failures = validate_extraction(metas, all_data_rows)
        failed_passes = [
            f"{failure.get('filename', '')}: {'; '.join(failure.get('issues', []))}"
            for failure in validation_failures
        ]
        batch = BatchResult(
            metas=metas, data_rows=all_data_rows,
            instruments=instruments, deduplicated_instruments=deduped,
            calibration_issues=cal_issues, conclusion_summary=conclusions,
            total_files=len(pdf_names), total_test_items=len(conclusions),
            extraction_quality="partial" if validation_failures else "complete",
            failed_passes=failed_passes,
            extraction_metrics={
                "pdf_count": len(pdf_names),
                "test_item_count": len(conclusions),
                "data_row_count": len(all_data_rows),
                "instrument_count": len(deduped),
                "validation_failure_count": len(validation_failures),
            },
        )
        structured_json = json_mod.dumps(_batch_result_to_dict(batch), ensure_ascii=False, default=str)

    return plain_text, structured_json


def _render_pdf_page_to_png(file_bytes: bytes, page_number: int = 0) -> bytes:
    """Render a single page of a PDF to PNG bytes via PyMuPDF.

    Used to feed PDF order forms into the visual (Qwen) extraction pipeline.
    """
    import fitz

    if not file_bytes.startswith(b"%PDF-"):
        raise ValueError("文件不是有效的 PDF")
    with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
        if page_number < 0 or page_number >= len(pdf):
            raise IndexError(f"PDF 页码超出范围: {page_number} (共 {len(pdf)} 页)")
        page = pdf[page_number]
        # Render at 2× scale for sufficient OCR quality; keep alpha off for
        # smaller payload size (visual extraction does not need transparency).
        pixmap = page.get_pixmap(matrix=fitz.Matrix(2.0, 2.0), alpha=False)
        return pixmap.tobytes("png")


def _infer_source_location(doc_type: str, field_name: str) -> str:
    """Infer a human-readable section label for an AI-extracted field."""
    if doc_type == "final_report":
        if any(kw in field_name for kw in ["报告编号","委托单位","设备名称","设备型号","制造商","检测依据"]):
            return "§1 报告概要"
        if any(kw in field_name for kw in ["样品","Sample","sample"]):
            return "§2 样品描述"
        if any(kw in field_name for kw in ["仪器","设备列表","Instrument"]):
            return "§3 检测设备"
        if any(kw in field_name for kw in ["结果","数据","Result","测试数据","CE","RE","RS","CS"]):
            return "§4 测试结果"
        if any(kw in field_name for kw in ["目录","TOC","toc"]):
            return "目录"
        if any(kw in field_name for kw in ["结论","conclusion","判定"]):
            return "§5 检测结论"
    elif doc_type == "test_plan":
        if any(kw in field_name for kw in ["委托","客户","单位","编号"]):
            return "§1 基本信息"
        if any(kw in field_name for kw in ["样品","设备","EUT","DUT"]):
            return "§2 被测设备信息"
        if any(kw in field_name for kw in ["测试项目","测试项","项目清单","标准"]):
            return "§3 测试项目清单"
        if any(kw in field_name for kw in ["细节","要求","布置","模式"]):
            return "§4 测试细节"
    return ""


async def _extract_and_update(set_id: str, doc_id: str, doc_type: str,
                               file_bytes: bytes, filename: str,
                               employee_id: str = "", doc_version: int = 1,
                               retry_count: int = 0):
    """Extract document content and update status in DB.

    Stores both plain_text for display and structured JSON for pipeline use.
    """
    try:
        t0 = time.time()
        base_meta = {
            "set_id": set_id,
            "doc_id": doc_id,
            "doc_type": doc_type,
            "filename": filename,
            "file_size_bytes": len(file_bytes),
            "file_size_kb": len(file_bytes) // 1024,
            "version": doc_version,
            "employee_id": employee_id,
            "retry_count": retry_count,
            "json_parse_retries": 0,
            "empty_response_retries": 0,
        }
        logger.info("document_extract_start", **base_meta)
        await database.update_doc_extraction_status_async(doc_id, "extracting")

        plain_text = ""
        structured_json = ""
        extraction_status = "done"
        extraction_quality = "complete"
        extraction_error_msg = ""
        extraction_meta: dict = {
            **base_meta,
        }

        if doc_type == "order_form":
            from services.order_form_extractor import OrderFormExtractor
            # Image-only forms use the unified Qwen vision gateway.
            from services.order_form_visual_extractor import OrderFormVisualExtractor
            if OrderFormVisualExtractor.is_image_extension(filename):
                data = await OrderFormVisualExtractor.extract(file_bytes)
                cell_map = {}       # image evidence is handled by the graph extraction stage
                cell_values = {}
            elif filename.lower().endswith(".pdf"):
                # PDF order forms: render page 1 as PNG → visual extraction.
                # The deterministic .xls parser cannot read PDFs, and scanning
                # a printed/scanned form as an image is the intended fallback.
                png_bytes = _render_pdf_page_to_png(file_bytes, page_number=0)
                data = await OrderFormVisualExtractor.extract(png_bytes)
                cell_map = {}
                cell_values = {}
            else:
                data = await OrderFormExtractor.extract(file_bytes)
                cell_map = OrderFormExtractor.get_cell_map()
                cell_values = OrderFormExtractor.get_all_cell_values(file_bytes)
            structured_json = data.model_dump_json()
            # Human-readable text for display — flatten nested structures
            dumped = data.model_dump()
            lines: list[str] = []
            for k, v in dumped.items():
                if not v:
                    continue
                if isinstance(v, dict):
                    # Flatten one level: "section: key=val | key=val ..."
                    sub = ", ".join(f"{sk}={sv}" for sk, sv in v.items() if sv and not isinstance(sv, (dict, list)))
                    if sub:
                        lines.append(f"{k}: {sub}")
                    # Handle nested contact dict
                    contact = v.get("contact")
                    if isinstance(contact, dict):
                        contact_str = ", ".join(f"{ck}={cv}" for ck, cv in contact.items() if cv)
                        if contact_str:
                            lines.append(f"{k}.contact: {contact_str}")
                elif isinstance(v, list):
                    lines.append(f"{k}: {', '.join(str(sv) for sv in v[:10])}")
                else:
                    lines.append(f"{k}: {v}")
            plain_text = "\n".join(lines)

        elif doc_type == "test_plan":
            from services.test_plan_extractor import TestPlanExtractor
            plain_text = extract_text(file_bytes, filename)
            try:
                data = await TestPlanExtractor.extract(file_bytes, filename)
                if data:
                    structured_json = data.model_dump_json()
                    metrics = getattr(data, "extraction_metrics", {}) or {}
                    if isinstance(metrics, dict):
                        extraction_meta.update(metrics)
                    failed_passes = getattr(data, "failed_passes", []) or []
                    if failed_passes:
                        extraction_status = "partial"
                        extraction_quality = "partial"
                        extraction_error_msg = "部分抽取失败: " + ", ".join(
                            str(p) for p in failed_passes[:10]
                        )
                        extraction_meta["failed_passes"] = list(failed_passes)
                else:
                    extraction_error_msg = "AI 提取返回空结果"
            except Exception as e:
                extraction_error_msg = str(e)
                logger.warning("test_plan_ai_extraction_failed",
                              doc_id=doc_id, filename=filename, error=extraction_error_msg)
            # If AI extraction produced no structured data, this is a failure
            if not structured_json:
                raise Exception(extraction_error_msg or "试验计划 AI 结构化提取失败")

        elif doc_type == "final_report":
            from services.report_extractor import ReportExtractor
            plain_text = extract_text(file_bytes, filename)
            try:
                data = await ReportExtractor.extract(file_bytes, filename)
                if data:
                    structured_json = data.model_dump_json()
                    metrics = getattr(data, "extraction_metrics", {}) or {}
                    if isinstance(metrics, dict):
                        extraction_meta.update(metrics)
                    failed_passes = getattr(data, "failed_passes", []) or []
                    if failed_passes:
                        extraction_status = "partial"
                        extraction_quality = "partial"
                        extraction_error_msg = "部分抽取失败: " + ", ".join(str(p) for p in failed_passes[:10])
                        extraction_meta["failed_passes"] = list(failed_passes)
                else:
                    extraction_error_msg = "AI 提取返回空结果"
            except Exception as e:
                extraction_error_msg = str(e)
                logger.warning("report_ai_extraction_failed",
                              doc_id=doc_id, filename=filename, error=extraction_error_msg)
            # If AI extraction produced no structured data, this is a failure
            if not structured_json:
                raise Exception(extraction_error_msg or "检测报告 AI 结构化提取失败")

        elif doc_type in ("test_standard",):
            plain_text = extract_text(file_bytes, filename)

        elif doc_type == "original_records":
            import asyncio
            plain_text = ""
            try:
                from services.document_unitizer import archive_member_manifest
                member_manifest = await asyncio.to_thread(
                    archive_member_manifest, file_bytes,
                )
                await database.replace_document_archive_members_async(
                    doc_id, member_manifest,
                )
                extraction_meta["archive_member_count"] = len(member_manifest)
                loop = asyncio.get_event_loop()
                plain_text, structured_json = await loop.run_in_executor(
                    None, _extract_raw_records_sync, file_bytes)
                if structured_json:
                    from services.raw_records_llm_extractor import enrich_raw_records_json
                    deepseek_user_id = "reviewer_" + hashlib.sha256(
                        (employee_id or "system").encode("utf-8")
                    ).hexdigest()[:24]
                    structured_json, semantic_failures = await enrich_raw_records_json(
                        structured_json,
                        user_id=deepseek_user_id,
                    )
                    extraction_meta["raw_record_semantic_failures"] = semantic_failures
                    if semantic_failures:
                        extraction_quality = "partial"
                    parsed_batch = json_mod.loads(structured_json)
                    existing_failures = list(parsed_batch.get("failed_passes", []) or [])
                    semantic_failure_text = [
                        f"语义提取失败: {failure}"
                        for failure in semantic_failures
                    ]
                    all_failures = existing_failures + semantic_failure_text
                    parsed_batch["failed_passes"] = all_failures
                    parsed_batch["extraction_quality"] = (
                        "partial" if all_failures else
                        parsed_batch.get("extraction_quality", "complete")
                    )
                    parsed_metrics = dict(parsed_batch.get("extraction_metrics", {}) or {})
                    parsed_metrics["semantic_failure_count"] = len(semantic_failures)
                    parsed_batch["extraction_metrics"] = parsed_metrics
                    structured_json = json_mod.dumps(parsed_batch, ensure_ascii=False)
                    if all_failures:
                        extraction_status = "partial"
                        extraction_quality = "partial"
                        extraction_error_msg = "部分抽取失败: " + ", ".join(
                            str(p) for p in all_failures[:10]
                        )
                        extraction_meta["failed_passes"] = all_failures
                if not structured_json:
                    extraction_error_msg = "原始记录 ZIP 中未找到可解析的 PDF 文件"
            except Exception as e:
                extraction_error_msg = str(e)
                logger.warning("raw_records_extraction_failed",
                              doc_id=doc_id, filename=filename, error=extraction_error_msg)
                # Fall back to plain text listing
                import zipfile, io
                try:
                    with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
                        names = zf.namelist()
                        plain_text = f"ZIP 包含 {len(names)} 个文件:\n" + "\n".join(names[:50])
                except Exception:
                    plain_text = extract_text(file_bytes, filename)
            if not structured_json:
                raise Exception(extraction_error_msg or "原始记录结构化提取失败")

        # Save structured extraction result BEFORE updating status to "done"
        # to avoid a race where the frontend polls "done" but metadata
        # hasn't been persisted yet.
        if structured_json:
            extraction_meta.setdefault("text_length", len(plain_text or ""))
            extraction_meta["duration_ms"] = round((time.time() - t0) * 1000)
            extraction_meta["quality"] = extraction_quality
            # Extract per-field _meta (confidence + source_quote) if AI provided it
            field_meta: dict[str, dict] = {}
            try:
                parsed = json_mod.loads(structured_json)
                meta_section = parsed.pop("_meta", None)
                if isinstance(meta_section, dict):
                    field_meta = {k: v for k, v in meta_section.items() if isinstance(v, dict)}
                # Re-serialize without _meta for the structured field
                if meta_section is not None:
                    structured_json = json_mod.dumps(parsed, ensure_ascii=False)
            except Exception:
                pass

            meta_fields = [
                {"field_name": "__structured__",
                 "field_value": structured_json,
                 "source_text": plain_text[:500],
                 "confidence": 1.0 if doc_type == "order_form" else 0.8},
                {"field_name": "__extraction_metrics__",
                 "field_value": json_mod.dumps(extraction_meta, ensure_ascii=False),
                 "source_text": "",
                 "confidence": 1.0},
            ]
            # Save per-field metadata from AI _meta
            for fname, finfo in field_meta.items():
                confidence = float(finfo.get("confidence", 0.8))
                source_quote = str(finfo.get("source_quote", ""))[:500]
                if confidence < 0 or confidence > 1:
                    confidence = 0.8
                meta_fields.append({
                    "field_name": fname,
                    "field_value": "",
                    "source_text": source_quote,
                    "confidence": confidence,
                    "source_location": _infer_source_location(doc_type, fname),
                })
            # For order_form, also save cell coordinate map and all cell values
            if doc_type == "order_form" and cell_map:
                meta_fields.append({
                    "field_name": "__cell_map__",
                    "field_value": json_mod.dumps(cell_map, ensure_ascii=False),
                    "source_text": "",
                    "confidence": 1.0,
                })
                meta_fields.append({
                    "field_name": "__cell_values__",
                    "field_value": json_mod.dumps(cell_values, ensure_ascii=False),
                    "source_text": "",
                    "confidence": 1.0,
                })
            await database.save_extracted_metadata_async(
                set_id, doc_id, meta_fields,
            )

        await database.update_doc_extraction_status_async(
            doc_id, extraction_status, plain_text,
            extraction_error=extraction_error_msg,
            extraction_quality=extraction_quality,
            extraction_meta=extraction_meta,
        )
        logger.info(
            "document_extract_done",
            **base_meta,
            status=extraction_status,
            extraction_quality=extraction_quality,
            duration_ms=extraction_meta.get("duration_ms", 0),
            text_length=extraction_meta.get("text_length", 0),
            warning_count=len(extraction_meta.get("failed_passes", []) or []),
        )

    except Exception as e:
        # Build a richer error message: include failure details from ExtractionError
        from services.exceptions import ExtractionError
        err_msg = str(e)
        if isinstance(e, ExtractionError) and e.failures:
            details = "; ".join(
                str(f.get("errors", f)) for f in e.failures if f
            )[:500]
            if details:
                err_msg = f"{err_msg} — 原因: {details}"
        logger.error("extraction_failed", set_id=set_id, doc_id=doc_id,
                     doc_type=doc_type, filename=filename, error=err_msg)
        failed_meta = {
            "doc_type": doc_type,
            "filename": filename,
            "file_size_bytes": len(file_bytes),
            "file_size_kb": len(file_bytes) // 1024,
            "version": doc_version,
            "employee_id": employee_id,
            "duration_ms": round((time.time() - t0) * 1000) if "t0" in locals() else 0,
            "quality": "failed",
            "retryable": doc_type in _AI_DOC_TYPES,
            "failure_stage": "extract_structure" if doc_type in _AI_DOC_TYPES else "extract",
            "retry_count": retry_count,
            "json_parse_retries": 0,
            "empty_response_retries": 0,
        }
        logger.error(
            "document_extract_failed",
            set_id=set_id,
            doc_id=doc_id,
            doc_type=doc_type,
            filename=filename,
            file_size_kb=len(file_bytes) // 1024,
            version=doc_version,
            employee_id=employee_id,
            duration_ms=failed_meta["duration_ms"],
            retryable=failed_meta["retryable"],
            extraction_stage=failed_meta["failure_stage"],
            error=err_msg,
        )
        await database.update_doc_extraction_status_async(
            doc_id, "failed", plain_text="", extraction_error=err_msg,
            extraction_quality="failed",
            extraction_meta=failed_meta,
        )


# Doc types that use AI APIs (DeepSeek) — these are gated by semaphore + timeout.
_AI_DOC_TYPES = {"test_plan", "final_report"}


async def _extract_with_timeout(set_id: str, doc_id: str, doc_type: str,
                                 file_bytes: bytes, filename: str,
                                 employee_id: str = "", doc_version: int = 1,
                                 retry_count: int = 0):
    """Wrap _extract_and_update with optional semaphore + timeout for AI extractions.

    Deterministic extractions (order_form) and CPU-bound extractions
    (original_records) skip the semaphore so they never queue behind slow AI calls.
    """
    task_key = f"{set_id}:{doc_id}"
    use_semaphore = doc_type in _AI_DOC_TYPES
    try:
        if use_semaphore:
            async with _extraction_semaphore:
                logger.info("extraction_start", set_id=set_id, doc_id=doc_id,
                           doc_type=doc_type, filename=filename)
                await asyncio.wait_for(
                    _extract_and_update(set_id, doc_id, doc_type, file_bytes, filename, employee_id, doc_version, retry_count),
                    timeout=EXTRACTION_TOTAL_TIMEOUT_SEC,
                )
        else:
            logger.info("extraction_start", set_id=set_id, doc_id=doc_id,
                       doc_type=doc_type, filename=filename)
            # Non-AI extractions still get a generous timeout (prevents thread-pool hangs)
            await asyncio.wait_for(
                _extract_and_update(set_id, doc_id, doc_type, file_bytes, filename, employee_id, doc_version, retry_count),
                timeout=EXTRACTION_TOTAL_TIMEOUT_SEC,
            )
    except asyncio.TimeoutError:
        err_msg = f"提取超时（超过 {EXTRACTION_TOTAL_TIMEOUT_SEC}s 限制），AI 服务可能繁忙"
        logger.error("extraction_timeout", set_id=set_id, doc_id=doc_id,
                    doc_type=doc_type, filename=filename,
                    timeout_sec=EXTRACTION_TOTAL_TIMEOUT_SEC)
        logger.error("document_extract_failed", set_id=set_id, doc_id=doc_id,
                    doc_type=doc_type, filename=filename,
                    file_size_kb=len(file_bytes) // 1024, version=doc_version,
                    employee_id=employee_id, duration_ms=EXTRACTION_TOTAL_TIMEOUT_SEC * 1000,
                    retryable=True, extraction_stage="timeout", error=err_msg)
        await database.update_doc_extraction_status_async(
            doc_id, "failed", plain_text="", extraction_error=err_msg,
            extraction_quality="failed",
            extraction_meta={
                "doc_type": doc_type,
                "filename": filename,
                "file_size_bytes": len(file_bytes),
                "file_size_kb": len(file_bytes) // 1024,
                "version": doc_version,
                "employee_id": employee_id,
                "quality": "failed",
                "failure_reason": "timeout",
                "timeout_sec": EXTRACTION_TOTAL_TIMEOUT_SEC,
                "duration_ms": EXTRACTION_TOTAL_TIMEOUT_SEC * 1000,
                "retry_count": retry_count,
            },
        )
    except asyncio.CancelledError:
        logger.info("extraction_cancelled", set_id=set_id, doc_id=doc_id,
                   doc_type=doc_type, filename=filename)
        await database.update_doc_extraction_status_async(
            doc_id, "failed", plain_text="", extraction_error="用户取消了提取任务",
            extraction_quality="failed",
            extraction_meta={
                "doc_type": doc_type,
                "filename": filename,
                "file_size_bytes": len(file_bytes),
                "file_size_kb": len(file_bytes) // 1024,
                "version": doc_version,
                "employee_id": employee_id,
                "quality": "failed",
                "failure_reason": "cancelled",
                "retry_count": retry_count,
            },
        )
    finally:
        _running_extractions.pop(task_key, None)


@router.post("/{set_id}/lock")
async def lock_set(set_id: str, request: Request = None, body: dict | None = Body(default=None),
                   user: User = Depends(require_role("admin", "reviewer"))):
    await _check_ownership(set_id, user)
    overview = await DocumentSetService.get_set_overview(set_id)
    if not overview:
        raise HTTPException(404, detail=f"文档集不存在: {set_id}")
    if not overview.get("is_ready"):
        missing = overview.get("missing_types", [])
        raise HTTPException(400, detail=f"文档不完整，缺少: {missing}")
    auto_gate_requested = bool((body or {}).get("auto_gate"))
    exception_only_requested = bool((body or {}).get("exception_only"))
    quality_gate = overview.get("extraction_gate") or {}
    if auto_gate_requested and quality_gate.get("status") != "pass":
        blockers = quality_gate.get("blockers") or []
        blocker_text = "、".join(
            f"{item.get('label') or item.get('doc_type')}: {item.get('reason')}"
            for item in blockers
        )
        logger.warning(
            "document_set_auto_gate_blocked", set_id=set_id,
            blocker_count=len(blockers), user=user.employee_id,
        )
        raise HTTPException(
            409,
            detail={
                "message": "自动质量门禁未通过，请处理异常资料后重试",
                "quality_gate": quality_gate,
                "blockers": blocker_text,
            },
        )
    required_exception_types = {
        str(item.get("doc_type"))
        for item in (quality_gate.get("blockers") or [])
        if item.get("doc_type")
    }
    unreviewed = [
        item.get("label") or item.get("doc_type")
        for item in overview.get("files", [])
        if item.get("required") and not item.get("reviewed_at")
        and (not exception_only_requested or item.get("doc_type") in required_exception_types)
    ]
    if unreviewed and not auto_gate_requested:
        raise HTTPException(409, detail="以下最新文档尚未完成核查：" + "、".join(unreviewed))
    standard_resolution = await database.get_document_set_standard_resolution_async(set_id)
    if standard_resolution["unresolved"]:
        codes = [item["reference_code"] for item in standard_resolution["unresolved"]]
        logger.warning(
            "document_set_standard_references_unresolved", set_id=set_id,
            reference_codes=codes, user=user.employee_id,
        )
        raise HTTPException(
            409,
            detail=(
                "文档中引用了尚未处理的标准：" + "、".join(codes) +
                "。请在标准选择中选用对应版本，或逐项填写跳过原因。"
            ),
        )
    try:
        ok = await DocumentSetService.lock_set(set_id)
        if not ok:
            raise HTTPException(400, detail="无法锁定: 状态不是 'incomplete' 或 'revision'")
        logger.info("lock_set", set_id=set_id, user=user.employee_id,
                     extra={"previous_status": overview.get("status"), "new_status": "locked",
                            "auto_gate": auto_gate_requested,
                            "exception_only": exception_only_requested})
        await save_audit_log_async(
            client_ip=_client_ip(request), action="set_lock",
            filename=set_id, employee_id=user.employee_id,
            detail=json_mod.dumps({"previous_status": overview.get("status"), "new_status": "locked"}))
        standard_review_enabled = bool(overview.get("standards"))
        if not standard_review_enabled:
            logger.info("standard_review_skipped", set_id=set_id,
                        reason="no_standard_selected", user=user.employee_id)
        return {
            "set_id": set_id,
            "status": "locked",
            "quality_gate": quality_gate,
            "extraction_review_skipped": auto_gate_requested,
            "extraction_review_mode": "exception_only" if exception_only_requested else "full",
            "standard_review_enabled": standard_review_enabled,
            "standard_review_notice": "" if standard_review_enabled else (
                "未选择测试标准，本次审核不会执行标准条款相关审查"
            ),
        }
    except HTTPException:
        raise
    except Exception as e:
        logger.error("lock_set_failed", set_id=set_id, user=user.employee_id, error=str(e))
        raise


@router.post("/{set_id}/revision")
async def create_revision(
    set_id: str,
    request: Request = None,
    user: User = Depends(require_role("admin", "reviewer")),
):
    """Reopen a reviewed document set as a revision draft.

    The previous review result remains visible until the draft is locked and
    reviewed again. File replacements are stored as new document versions.
    """
    await _check_ownership(set_id, user)
    overview = await DocumentSetService.get_set_overview(set_id)
    if not overview:
        raise HTTPException(404, detail=f"文档集不存在: {set_id}")
    if overview.get("status") == "revision":
        return {"set_id": set_id, "status": "revision"}
    if overview.get("status") != "reviewed":
        raise HTTPException(400, detail="只有已审核的文档集可以创建修订")

    try:
        ok = await DocumentSetService.create_revision(set_id)
        if not ok:
            raise HTTPException(400, detail="创建修订失败")
        logger.info("create_revision", set_id=set_id, user=user.employee_id,
                    extra={"previous_status": "reviewed", "new_status": "revision"})
        await save_audit_log_async(
            client_ip=_client_ip(request), action="set_revision_create",
            filename=set_id, employee_id=user.employee_id,
            detail=json_mod.dumps({"previous_status": "reviewed", "new_status": "revision"}))
        return {"set_id": set_id, "status": "revision"}
    except HTTPException:
        raise
    except Exception as e:
        logger.error("create_revision_failed", set_id=set_id, user=user.employee_id, error=str(e))
        raise HTTPException(500, detail=f"创建修订失败: {str(e)}")


@router.put("/{set_id}/status")
async def update_status(
    set_id: str, body: dict,
    request: Request = None,
    user: User = Depends(require_role("admin", "reviewer")),
):
    await _check_ownership(set_id, user)
    new_status = body.get("status", "")
    try:
        # Capture old status before update
        old_ds = await database.get_document_set_async(set_id)
        old_status = old_ds["status"] if old_ds else "unknown"
        allowed_transitions = {
            "revision": ("reviewed", "locked", "reviewing", "error"),
            "error": ("reviewing",),
        }
        allowed_from = allowed_transitions.get(new_status)
        if not allowed_from:
            raise HTTPException(400, detail=f"不允许通过此接口设置状态: {new_status}")
        ok = await database.transition_set_status_async(set_id, new_status, allowed_from)
        if not ok:
            raise HTTPException(409, detail=f"状态迁移冲突: {old_status} -> {new_status}")
        logger.info("update_status", set_id=set_id, user=user.employee_id,
                     extra={"old_status": old_status, "new_status": new_status})
        await save_audit_log_async(
            client_ip=_client_ip(request), action="set_status_change",
            filename=set_id, employee_id=user.employee_id,
            detail=json_mod.dumps({"old_status": old_status, "new_status": new_status}))
        return {"set_id": set_id, "status": new_status}
    except HTTPException:
        raise
    except Exception as e:
        logger.error("update_status_failed", set_id=set_id, user=user.employee_id,
                      error=str(e), extra={"new_status": new_status})
        raise


@router.get("/{set_id}/versions/{doc_type}")
async def get_version_history(
    set_id: str, doc_type: DocType,
    user: User = Depends(require_role("admin", "reviewer", "viewer")),
):
    await _check_ownership(set_id, user)
    if doc_type not in DOC_TYPES:
        raise HTTPException(422, detail=f"无效的文档类型: {doc_type}")
    versions = await DocumentSetService.get_version_history(set_id, doc_type)
    return {"set_id": set_id, "doc_type": doc_type,
            "versions": [v.model_dump() for v in versions], "total": len(versions)}


@router.get("/{set_id}/diff")
async def diff_versions(
    set_id: str,
    old_doc_id: str = Query(...), new_doc_id: str = Query(...),
    user: User = Depends(require_role("admin", "reviewer", "viewer")),
):
    await _check_ownership(set_id, user)
    await _require_document_in_set(set_id, old_doc_id)
    await _require_document_in_set(set_id, new_doc_id)
    diffs = await DocumentSetService.diff_versions(old_doc_id, new_doc_id)
    changed = [d.model_dump() for d in diffs if d.changed]
    return {"old_doc_id": old_doc_id, "new_doc_id": new_doc_id,
            "total_fields": len(diffs), "changed_fields": len(changed),
            "diffs": [d.model_dump() for d in diffs], "changed": changed}



@router.get("/{set_id}/extractions")
async def get_extractions(
    set_id: str, user: User = Depends(require_role("admin", "reviewer", "viewer")),
):
    """Get structured extraction data for all documents in a set.

    Returns parsed extraction results keyed by doc_type, ready for review modals.
    """
    ds = await database.get_document_set_async(set_id)
    if not ds:
        raise HTTPException(404, detail=f"文档集不存在: {set_id}")
    if user.role != "admin" and ds.get("employee_id", "") != user.employee_id:
        raise HTTPException(403, detail="无权操作此文档集")

    extractions: dict[str, dict] = {}
    # Group documents by type and pick the latest version for each
    docs_by_type: dict[str, list[dict]] = {}
    for doc in ds.get("documents", []):
        dt = doc["doc_type"]
        docs_by_type.setdefault(dt, []).append(doc)
    # Sort each group by doc_version descending so we pick the latest
    for dt in docs_by_type:
        docs_by_type[dt].sort(key=lambda d: d.get("doc_version", 1), reverse=True)

    for dt, docs in docs_by_type.items():
        doc = docs[0]  # latest version
        if doc.get("extraction_status") not in {"done", "partial"}:
            extractions[dt] = {"status": doc.get("extraction_status", "pending"),
                                "doc_id": doc["doc_id"],
                                "filename": doc.get("filename", ""),
                                "plain_text": doc.get("plain_text", ""),
                                "error": doc.get("extraction_error", ""),
                                "extraction_quality": doc.get("extraction_quality", ""),
                                "extraction_meta": doc.get("extraction_meta", {})}
            continue

        result: dict = {"status": doc.get("extraction_status", "done"), "doc_id": doc["doc_id"],
                        "filename": doc.get("filename", ""),
                        "plain_text": doc.get("plain_text", ""),
                        "error": doc.get("extraction_error", ""),
                        "extraction_quality": doc.get("extraction_quality", ""),
                        "extraction_meta": doc.get("extraction_meta", {}),
                        "fields": [], "structured": None,
                        "cell_map": None, "cell_values": None}

        meta_rows = await database.get_extracted_metadata_async(doc["doc_id"])
        for row in meta_rows:
            if row.get("field_name") == "__structured__":
                try:
                    result["structured"] = json_mod.loads(row["field_value"])
                except Exception:
                    pass
            elif row.get("field_name") == "__cell_map__":
                try:
                    result["cell_map"] = json_mod.loads(row["field_value"])
                except Exception:
                    pass
            elif row.get("field_name") == "__cell_values__":
                try:
                    result["cell_values"] = json_mod.loads(row["field_value"])
                except Exception:
                    pass
            elif row.get("field_name") == "__extraction_metrics__":
                try:
                    result["extraction_meta"] = json_mod.loads(row["field_value"])
                except Exception:
                    pass
            else:
                result["fields"].append({
                    "field_name": row["field_name"],
                    "field_value": row["field_value"],
                    "source_text": row.get("source_text", ""),
                    "confidence": row.get("confidence", 1.0),
                    "source_location": row.get("source_location", ""),
                })

        extractions[dt] = result

    return {"set_id": set_id, "extractions": extractions}


@router.get("/{set_id}/documents/{doc_id}/extraction")
async def get_document_extraction(
    set_id: str, doc_id: str,
    user: User = Depends(require_role("admin", "reviewer", "viewer")),
):
    """Get structured extraction data for a single document.

    This is the preferred endpoint for review modals — it only loads one
    document's metadata, decoupling it from other documents' states.
    """
    ds = await database.get_document_set_async(set_id)
    if not ds:
        raise HTTPException(404, detail=f"文档集不存在: {set_id}")
    if user.role != "admin" and ds.get("employee_id", "") != user.employee_id:
        raise HTTPException(403, detail="无权操作此文档集")

    # Find the target document
    target = None
    for doc in ds.get("documents", []):
        if doc["doc_id"] == doc_id:
            target = doc
            break
    if not target:
        raise HTTPException(404, detail=f"文档不存在: {doc_id}")

    doc_type = target["doc_type"]
    if target.get("extraction_status") not in {"done", "partial"}:
        return {
            "set_id": set_id, "doc_id": doc_id, "doc_type": doc_type,
            "status": target.get("extraction_status", "pending"),
            "filename": target.get("filename", ""),
            "plain_text": target.get("plain_text", ""),
            "extraction_quality": target.get("extraction_quality", ""),
            "extraction_meta": target.get("extraction_meta", {}),
            "fields": [], "structured": None,
            "cell_map": None, "cell_values": None,
            "error": target.get("extraction_error", ""),
        }

    result: dict = {
        "set_id": set_id, "doc_id": doc_id, "doc_type": doc_type,
        "status": target.get("extraction_status", "done"),
        "filename": target.get("filename", ""),
        "plain_text": target.get("plain_text", ""),
        "extraction_quality": target.get("extraction_quality", ""),
        "extraction_meta": target.get("extraction_meta", {}),
        "error": target.get("extraction_error", ""),
        "fields": [], "structured": None,
        "cell_map": None, "cell_values": None,
    }

    meta_rows = await database.get_extracted_metadata_async(doc_id)
    for row in meta_rows:
        if row.get("field_name") == "__structured__":
            try:
                result["structured"] = json_mod.loads(row["field_value"])
            except (json_mod.JSONDecodeError, TypeError) as e:
                logger.warning("structured_json_parse_failed", doc_id=doc_id,
                              set_id=set_id, error=str(e))
        elif row.get("field_name") == "__cell_map__":
            try:
                result["cell_map"] = json_mod.loads(row["field_value"])
            except (json_mod.JSONDecodeError, TypeError) as e:
                logger.warning("cell_map_json_parse_failed", doc_id=doc_id,
                              set_id=set_id, error=str(e))
        elif row.get("field_name") == "__cell_values__":
            try:
                result["cell_values"] = json_mod.loads(row["field_value"])
            except (json_mod.JSONDecodeError, TypeError) as e:
                logger.warning("cell_values_json_parse_failed", doc_id=doc_id,
                              set_id=set_id, error=str(e))
        elif row.get("field_name") == "__extraction_metrics__":
            try:
                result["extraction_meta"] = json_mod.loads(row["field_value"])
            except (json_mod.JSONDecodeError, TypeError) as e:
                logger.warning("extraction_metrics_json_parse_failed", doc_id=doc_id,
                              set_id=set_id, error=str(e))
        else:
            result["fields"].append({
                "field_name": row["field_name"],
                "field_value": row["field_value"],
                "source_text": row.get("source_text", ""),
                "confidence": row.get("confidence", 1.0),
                "source_location": row.get("source_location", ""),
            })

    return result


@router.patch("/{set_id}/documents/{doc_id}/overrides")
async def save_overrides(
    set_id: str, doc_id: str,
    body: dict = Body(...),
    user: User = Depends(require_role("admin", "reviewer")),
):
    """Save human-edited field values for a document."""
    await _check_ownership(set_id, user)
    await _require_editable_set(set_id)
    await _require_document_in_set(set_id, doc_id)
    overrides: dict = body.get("overrides", {})
    if not isinstance(overrides, dict):
        raise HTTPException(400, "overrides must be a dict of field_name -> value")
    try:
        count = await database.save_human_overrides_async(doc_id, overrides)
        if overrides:
            await database.clear_document_review_async(doc_id)
        logger.info("human_overrides_saved", set_id=set_id, doc_id=doc_id,
                    field_count=count, user=user.employee_id)
        return {"saved": count}
    except Exception as e:
        logger.error("save_overrides_failed", set_id=set_id, doc_id=doc_id,
                     error=str(e))
        raise HTTPException(500, f"保存失败: {str(e)}")


@router.post("/{set_id}/documents/{doc_id}/review-confirmation")
async def confirm_document_review(
    set_id: str, doc_id: str,
    request: Request = None,
    user: User = Depends(require_role("admin", "reviewer")),
):
    """Persist the reviewer and time for the current document version."""
    await _check_ownership(set_id, user)
    await _require_editable_set(set_id)
    await _require_document_in_set(set_id, doc_id)
    try:
        result = await database.confirm_document_review_async(
            set_id, doc_id, user.employee_id,
        )
    except ValueError as exc:
        raise HTTPException(409, detail=str(exc)) from exc
    await save_audit_log_async(
        client_ip=_client_ip(request), action="document_extraction_review_confirmed",
        filename=doc_id, employee_id=user.employee_id,
        detail=json_mod.dumps({"set_id": set_id, "doc_id": doc_id}),
    )
    return result


@router.get("/{set_id}/documents/{doc_id}/overrides")
async def get_overrides(
    set_id: str, doc_id: str,
    user: User = Depends(require_role("admin", "reviewer", "viewer")),
):
    """Get human override values for a document."""
    await _check_ownership(set_id, user)
    await _require_document_in_set(set_id, doc_id)
    overrides = await database.get_human_overrides_async(doc_id)
    return {"overrides": overrides, "count": len(overrides)}


@router.post("/{set_id}/documents/{doc_id}/retry-extraction")
async def retry_extraction(
    set_id: str, doc_id: str,
    request: Request = None,
    user: User = Depends(require_role("admin", "reviewer")),
):
    """Re-trigger extraction for a stuck or failed document.

    Recovers from lost async tasks (due to server restart) or transient AI failures.
    """
    await _check_ownership(set_id, user)
    ds = await _require_editable_set(set_id)

    # Find the target document
    target = None
    for doc in ds.get("documents", []):
        if doc.get("doc_id") == doc_id:
            target = doc
            break
    if not target:
        raise HTTPException(404, detail=f"文档不存在: {doc_id}")

    doc_type = target["doc_type"]
    filename = target.get("filename", "")

    # Load file content (required for extraction)
    import asyncio
    file_bytes = await database.get_file_content_async(doc_id)
    if not file_bytes:
        raise HTTPException(400, detail="原始文件已丢失，无法重新提取。请重新上传文档。")

    logger.info("retry_extraction", set_id=set_id, doc_id=doc_id,
               doc_type=doc_type, user=user.employee_id)
    retry_count = int((target.get("extraction_meta") or {}).get("retry_count") or 0) + 1
    logger.info(
        "document_extract_retry",
        set_id=set_id,
        doc_id=doc_id,
        doc_type=doc_type,
        filename=filename,
        version=int(target.get("doc_version") or 1),
        employee_id=user.employee_id,
        previous_error=target.get("extraction_error", "")[:300],
        retry_count=retry_count,
        next_action="background_extract",
    )

    # Re-trigger extraction in background
    task = asyncio.create_task(_extract_with_timeout(
        set_id, doc_id, doc_type, file_bytes, filename,
        user.employee_id, int(target.get("doc_version") or 1), retry_count,
    ))
    _running_extractions[f"{set_id}:{doc_id}"] = task

    return {"doc_id": doc_id, "status": "retrying", "message": f"{doc_type} 重新提取已触发"}


@router.post("/{set_id}/documents/{doc_id}/cancel-extraction")
async def cancel_extraction(
    set_id: str, doc_id: str,
    request: Request = None,
    user: User = Depends(require_role("admin", "reviewer")),
):
    """Cancel a running extraction task for a document.

    Useful when an AI extraction is stuck or the user uploaded the wrong file.
    """
    await _check_ownership(set_id, user)
    await _require_editable_set(set_id)
    task_key = f"{set_id}:{doc_id}"
    task = _running_extractions.get(task_key)
    if not task:
        # Check if doc is actually extracting
        ds = await database.get_document_set_async(set_id)
        if ds:
            for doc in ds.get("documents", []):
                if doc.get("doc_id") == doc_id:
                    if doc.get("extraction_status") == "extracting":
                        # Task reference lost (e.g. server restart) — force mark as failed
                        await database.update_doc_extraction_status_async(
                            doc_id, "failed", plain_text="",
                            extraction_error="提取任务丢失（服务器可能重启），请重新上传或重试")
                        return {"cancelled": True, "doc_id": doc_id,
                                "message": "提取任务已丢失，已标记为失败。请重试提取。"}
                    else:
                        return {"cancelled": False, "doc_id": doc_id,
                                "message": f"文档不在提取中（当前状态: {doc.get('extraction_status')}）"}
        return {"cancelled": False, "doc_id": doc_id, "message": "未找到运行中的提取任务"}

    task.cancel()
    logger.info("extraction_cancel_requested", set_id=set_id, doc_id=doc_id,
               user=user.employee_id)
    return {"cancelled": True, "doc_id": doc_id, "message": "提取任务已取消"}


@router.get("/{set_id}/documents/{doc_id}/file")
async def download_document_file(
    set_id: str, doc_id: str,
    user: User = Depends(require_role("admin", "reviewer", "viewer")),
):
    """Download the original uploaded file for a document."""
    await _check_ownership(set_id, user)
    await _require_document_in_set(set_id, doc_id)
    file_bytes = await database.get_file_content_async(doc_id)
    if file_bytes is None:
        raise HTTPException(404, detail="原始文件不存在或已被清理")

    # Get filename from document record for Content-Disposition
    ds = await database.get_document_set_async(set_id)
    filename = "download"
    if ds:
        for doc in ds.get("documents", []):
            if doc.get("doc_id") == doc_id:
                filename = doc.get("filename", "download")
                break

    # Determine media type from extension
    import mimetypes
    media_type, _ = mimetypes.guess_type(filename)
    if not media_type:
        media_type = "application/octet-stream"

    # Encode filename for Content-Disposition (RFC 5987)
    from urllib.parse import quote
    encoded_filename = quote(filename, safe="")

    return Response(
        content=file_bytes,
        media_type=media_type,
        headers={
            "Content-Disposition": f"inline; filename*=UTF-8''{encoded_filename}",
            "Content-Length": str(len(file_bytes)),
        },
    )


@router.get("/{set_id}/documents/{doc_id}/archive-members/{member_index}/file")
async def view_archive_member_file(
    set_id: str, doc_id: str, member_index: int,
    user: User = Depends(require_role("admin", "reviewer", "viewer")),
):
    """Open one supported source file inside a raw-record ZIP.

    ``member_index`` follows the same one-based supported-file inventory used
    by document unitization. The archive path is never accepted from the
    client, avoiding path traversal and ambiguous duplicate names.
    """
    await _check_ownership(set_id, user)
    await _require_document_in_set(set_id, doc_id)
    ds = await database.get_document_set_async(set_id)
    target = next(
        (doc for doc in (ds or {}).get("documents", []) if doc.get("doc_id") == doc_id),
        None,
    )
    if not target or target.get("doc_type") != "original_records":
        raise HTTPException(422, detail="仅原始记录压缩包支持打开内部文件")
    file_bytes = await database.get_file_content_async(doc_id)
    if file_bytes is None:
        raise HTTPException(404, detail="原始文件不存在或已被清理")

    import mimetypes
    from services.document_preview import extract_archive_member
    try:
        display_filename, content = extract_archive_member(file_bytes, member_index)
        filename = Path(display_filename).name or f"record-{member_index}"
    except IndexError:
        raise HTTPException(404, detail="压缩包内文件不存在")
    except (zipfile.BadZipFile, RuntimeError, OSError, ValueError):
        raise HTTPException(422, detail="原始记录压缩包无法读取")

    media_type, _ = mimetypes.guess_type(filename)
    encoded_filename = quote(filename, safe="")
    return Response(
        content=content,
        media_type=media_type or "application/octet-stream",
        headers={
            "Content-Disposition": f"inline; filename*=UTF-8''{encoded_filename}",
            "Content-Length": str(len(content)),
        },
    )


@router.get("/{set_id}/documents/{doc_id}/archive-members")
async def list_archive_member_files(
    set_id: str, doc_id: str,
    user: User = Depends(require_role("admin", "reviewer", "viewer")),
):
    """List reviewed raw-record metadata with its real ZIP member index.

    Extraction sorts filenames for deterministic processing while document
    unitization keeps archive order. Joining by the stored source filename is
    therefore required; using the extraction array position can open the
    wrong source document.
    """
    await _check_ownership(set_id, user)
    await _require_document_in_set(set_id, doc_id)
    ds = await database.get_document_set_async(set_id)
    target = next(
        (doc for doc in (ds or {}).get("documents", []) if doc.get("doc_id") == doc_id),
        None,
    )
    if not target or target.get("doc_type") != "original_records":
        raise HTTPException(422, detail="仅原始记录压缩包支持文件清单")
    file_bytes = await database.get_file_content_async(doc_id)
    if file_bytes is None:
        raise HTTPException(404, detail="原始文件不存在或已被清理")

    import zipfile
    from services.document_unitizer import archive_member_manifest
    try:
        manifest = await database.get_document_archive_members_async(doc_id)
        if not manifest:
            manifest = await asyncio.to_thread(archive_member_manifest, file_bytes)
            await database.replace_document_archive_members_async(doc_id, manifest)
    except (zipfile.BadZipFile, RuntimeError, OSError, ValueError):
        raise HTTPException(422, detail="原始记录压缩包无法读取")

    indices_by_name: dict[str, list[int]] = {}
    archive_names: dict[int, str] = {}
    display_names: dict[int, str] = {}
    hashes_by_index: dict[int, str] = {}
    for member in manifest:
        index = int(member.get("member_index") or 0)
        raw_name = str(member.get("raw_filename") or "")
        display_name = str(member.get("display_filename") or raw_name)
        raw_basename = raw_name.replace("\\", "/").rsplit("/", 1)[-1]
        display_basename = display_name.replace("\\", "/").rsplit("/", 1)[-1]
        for candidate in {raw_basename, display_basename}:
            indices_by_name.setdefault(candidate, []).append(index)
        archive_names[index] = raw_basename
        display_names[index] = display_basename
        hashes_by_index[index] = str(member.get("content_hash") or "")

    structured: dict = {}
    for row in await database.get_extracted_metadata_async(doc_id):
        if row.get("field_name") != "__structured__":
            continue
        try:
            structured = json_mod.loads(row.get("field_value") or "{}")
        except (json_mod.JSONDecodeError, TypeError):
            structured = {}
        break

    members: list[dict] = []
    for meta in structured.get("metas", []) if isinstance(structured, dict) else []:
        if not isinstance(meta, dict):
            continue
        source_filename = str(meta.get("filename") or "")
        member_index = int(meta.get("archive_member_index") or 0)
        member_hash = str(meta.get("archive_member_hash") or "")
        if member_index not in archive_names or (
            member_hash and hashes_by_index.get(member_index) != member_hash
        ):
            member_index = 0
        if not member_index and member_hash:
            matches = [
                index for index, content_hash in hashes_by_index.items()
                if content_hash == member_hash
            ]
            member_index = matches[0] if len(matches) == 1 else 0
        if not member_index:
            candidates = indices_by_name.get(source_filename, [])
            member_index = candidates.pop(0) if candidates else 0
        header = meta.get("_header_fields") if isinstance(meta.get("_header_fields"), dict) else {}
        members.append({
            "member_index": member_index,
            "source_filename": source_filename,
            "filename": display_names.get(member_index) or source_filename,
            "test_item_name": str(meta.get("test_item_name") or meta.get("test_item_code") or ""),
            "sequence": str(meta.get("sequence") or ""),
            "test_mode": str(meta.get("test_mode") or ""),
            "sample_id": str(header.get("sample_id") or ""),
            "conclusion": str(header.get("test_conclusion") or ""),
        })

    # Preserve traceability even for files whose structured extraction failed.
    mapped = {item["member_index"] for item in members if item["member_index"]}
    for index, source_filename in archive_names.items():
        if index not in mapped:
            members.append({
                "member_index": index, "source_filename": source_filename,
                "filename": display_names.get(index) or source_filename,
                "test_item_name": "", "sequence": "", "test_mode": "",
                "sample_id": "", "conclusion": "",
            })
    members.sort(key=lambda item: item["member_index"] or 10 ** 9)
    return {
        "doc_id": doc_id,
        "filename": target.get("filename", ""),
        "total": len(manifest),
        "members": members,
    }
