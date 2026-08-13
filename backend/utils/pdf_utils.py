"""Shared PDF/DOCX/XLSX/XLS text extraction. Consolidates duplicate implementations."""

import io
import pypdf
from openpyxl import load_workbook
import xlrd


def extract_pdf_text(file_bytes: bytes) -> str:
    reader = pypdf.PdfReader(io.BytesIO(file_bytes))
    return "\n".join(t for p in reader.pages if (t := p.extract_text()))


def extract_docx_text(file_bytes: bytes) -> str:
    """Extract the text a reviewer can actually see in the rendered DOCX.

    Reading ``word/document.xml`` directly can include stale template runs or
    glyphs that Word/LibreOffice does not render.  Those invisible characters
    must not become audit evidence or change a verdict.
    """
    from services.document_unitizer import _office_to_pdf

    return extract_pdf_text(_office_to_pdf(file_bytes, "document.docx"))


def extract_xlsx_text(file_bytes: bytes) -> str:
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    chunks: list[str] = []
    for ws in wb.worksheets:
        chunks.append(f"## {ws.title}")
        for row in ws.iter_rows(values_only=True):
            values = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if values:
                chunks.append("\t".join(values))
    return "\n".join(chunks)


def extract_xls_text(file_bytes: bytes) -> str:
    wb = xlrd.open_workbook(file_contents=file_bytes)
    chunks: list[str] = []
    for ws in wb.sheets():
        chunks.append(f"## {ws.name}")
        for row_idx in range(ws.nrows):
            values = [
                str(ws.cell_value(row_idx, col_idx)).strip()
                for col_idx in range(ws.ncols)
                if ws.cell_type(row_idx, col_idx) != xlrd.XL_CELL_EMPTY
            ]
            non_empty = [v for v in values if v]
            if non_empty:
                chunks.append("\t".join(non_empty))
    return "\n".join(chunks)


def extract_text(file_bytes: bytes, filename: str) -> str:
    lower = filename.lower()
    if lower.endswith(".pdf"): return extract_pdf_text(file_bytes)
    elif lower.endswith(".docx"): return extract_docx_text(file_bytes)
    elif lower.endswith(".xlsx"): return extract_xlsx_text(file_bytes)
    elif lower.endswith(".xls"): return extract_xls_text(file_bytes)
    raise ValueError(f"不支持的文件格式: {filename}")
