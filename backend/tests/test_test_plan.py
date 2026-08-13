"""Tests for test plan extraction."""

from __future__ import annotations

import pytest
from pathlib import Path
from services.test_plan_extractor import (
    extract_text, count_test_codes, validate_extraction, ExtractionError,
    enforce_minimum_content, TestPlanExtractor, _link_test_details_to_items,
    _recover_tabular_sample_requirements,
    _merge_inventory_items, _page_aware_text,
    _transcribe_scan_pages,
    _repair_source_evidence,
)
from models import TestPlanData, TestPlanBasicInfo, TestPlanItem, TestPlanDetail

SAMPLE = Path(__file__).parent.parent.parent / "example" / "EMC测试大纲模板（2023版1204）-P57A2车标零件(1).pdf"


@pytest.fixture
def sample_bytes() -> bytes:
    if not SAMPLE.exists():
        pytest.skip("Sample not found")
    return SAMPLE.read_bytes()


def test_extract_text(sample_bytes: bytes):
    text = extract_text(sample_bytes, "test.pdf")
    assert len(text) > 1000
    assert "EMC" in text


def test_extract_xlsx_text_for_llm_input():
    """Excel plans are converted to plain text only; LLM still owns structuring."""
    from openpyxl import Workbook
    import io

    wb = Workbook()
    ws = wb.active
    ws.title = "DVP 内容"
    ws["A1"] = "Design Verification Plan"
    ws["B2"] = "测试项目"
    ws["C2"] = "EQ/MC01 传导发射"
    buf = io.BytesIO()
    wb.save(buf)

    text = extract_text(buf.getvalue(), "plan.xlsx")
    assert "## DVP 内容" in text
    assert "Design Verification Plan" in text
    assert "EQ/MC01 传导发射" in text


def test_count_codes(sample_bytes: bytes):
    text = extract_text(sample_bytes, "test.pdf")
    codes = count_test_codes(text)
    assert len(codes) >= 20
    assert "EQ/MC03" in codes


def test_count_codes_empty():
    assert count_test_codes("no codes") == set()


def test_validate_ok():
    text = "EQ/MC01 EQ/MC02"
    data = TestPlanData(
        basic_info=TestPlanBasicInfo(part_name="x"),
        test_items=[TestPlanItem(code="EQ/MC01"), TestPlanItem(code="EQ/MC02")],
        test_details=[TestPlanDetail(code="EQ/MC01"), TestPlanDetail(code="EQ/MC02")],
    )
    assert validate_extraction(text, data) == []


def test_validate_missing():
    text = "EQ/MC01 EQ/MC02"
    data = TestPlanData(test_items=[TestPlanItem(code="EQ/MC01")])
    assert len(validate_extraction(text, data)) >= 1


def test_validate_extra():
    text = "EQ/MC01"
    data = TestPlanData(test_items=[TestPlanItem(code="EQ/MC01"), TestPlanItem(code="EQ/MC99")])
    assert len(validate_extraction(text, data)) >= 1


def test_validate_missing_detail():
    """Missing test_details should NOT be an error — only a logged warning."""
    text = "EQ/MC01"
    data = TestPlanData(
        basic_info=TestPlanBasicInfo(part_name="x"),
        test_items=[TestPlanItem(code="EQ/MC01")],
        test_details=[],
    )
    assert validate_extraction(text, data) == []


def test_validate_missing_detail_subset():
    """Extraction succeeds when test_details are a subset of test_items.

    Reproduces the EMC-20260622-fv3iru scenario: 8 test items in the doc,
    but AI only extracts 2 test details. This should succeed because
    downstream checks handle missing/empty details gracefully.
    """
    text = "EQ/IC05 EQ/IR08 EQ/MC02 EQ/MC03 EQ/MR06 EQ/MR07 EQ/IC10 EQ/IC20"
    data = TestPlanData(
        basic_info=TestPlanBasicInfo(part_name="x"),
        test_items=[
            TestPlanItem(code="EQ/IC05"), TestPlanItem(code="EQ/IR08"),
            TestPlanItem(code="EQ/MC02"), TestPlanItem(code="EQ/MC03"),
            TestPlanItem(code="EQ/MR06"), TestPlanItem(code="EQ/MR07"),
            TestPlanItem(code="EQ/IC10"), TestPlanItem(code="EQ/IC20"),
        ],
        test_details=[
            TestPlanDetail(code="EQ/IC05"),
            TestPlanDetail(code="EQ/IR08"),
        ],
    )
    assert validate_extraction(text, data) == []
    # test_items checks still work: missing test_items IS still an error
    data2 = TestPlanData(
        basic_info=TestPlanBasicInfo(part_name="x"),
        test_items=[TestPlanItem(code="EQ/IC05")],  # only 1 of 8
        test_details=[TestPlanDetail(code="EQ/IC05")],
    )
    assert len(validate_extraction(text, data2)) >= 1  # missing test_items


def test_validate_ignores_empty_item_codes_for_name_only_plans():
    """Vendor plans may list semantic test names without EQ codes."""
    text = "验证计划 序号 试验项目 试验描述 接受标准 瞬态抗扰度试验 系统12V电源电压波动试验"
    data = TestPlanData(
        basic_info=TestPlanBasicInfo(part_name="灯具"),
        test_items=[
            TestPlanItem(code="", name="瞬态抗扰度试验"),
            TestPlanItem(code="", name="系统12V电源电压波动试验"),
        ],
        test_details=[],
    )

    assert validate_extraction(text, data) == []


def test_arbitrary_prose_plan_requires_source_anchored_llm_items():
    text = (
        "--- 文件第 3 页 ---\n本次需要进行系统12V电源电压波动试验，"
        "试验脉冲P4，测试后产品应能正常工作。"
    )
    data = TestPlanData(test_items=[TestPlanItem(
        code="系统12V电源电压波动试验",
        name="系统12V电源电压波动试验",
        source_quote="系统12V电源电压波动试验，试验脉冲P4",
        source_location="文件第3页",
    )])

    assert validate_extraction(
        text, data, require_source_evidence=True,
    ) == []
    data.test_items[0].source_quote = "文档中不存在的固定模板字段"
    assert "无法回查" in validate_extraction(
        text, data, require_source_evidence=True,
    )[0]


def test_inventory_audit_merges_name_only_item_without_template_code():
    primary = [TestPlanItem(code="EQ/IC01", name="瞬态抗扰度试验")]
    audit = [
        TestPlanItem(code="4.1 EQ/IC01", name="瞬态抗扰度试验 EQ/IC01"),
        TestPlanItem(code="系统12V电源电压波动试验", name="系统12V电源电压波动试验"),
    ]

    merged, added = _merge_inventory_items(primary, audit)

    assert added == 1
    assert [item.name for item in merged] == [
        "瞬态抗扰度试验", "系统12V电源电压波动试验",
    ]


def test_pdf_llm_input_contains_physical_page_markers(sample_bytes: bytes):
    text, _ = _page_aware_text(sample_bytes, "plan.pdf")

    assert "--- 文件第 1 页 ---" in text
    assert "--- 文件第 2 页 ---" in text


@pytest.mark.asyncio
async def test_scan_plan_page_is_visually_transcribed_in_overlapping_tiles():
    import fitz

    document = fitz.open()
    document.new_page()
    pdf_bytes = document.tobytes()
    document.close()

    class FakeGateway:
        def __init__(self):
            self.calls = 0

        async def call_json(self, task, **kwargs):
            self.calls += 1
            return {
                "page_text": f"区域{self.calls}测试项目原文",
                "coverage_state": "complete",
            }

    gateway = FakeGateway()
    result = await _transcribe_scan_pages(
        pdf_bytes, "任意版式计划.pdf", [(1, "")], gateway=gateway,
    )

    assert gateway.calls == 2
    assert result[1] == "区域1测试项目原文\n区域2测试项目原文"


def test_non_verbatim_llm_citation_is_repaired_from_declared_page():
    raw = (
        "--- 文件第 8 页 ---\n其他项目\n"
        "--- 文件第 9 页 ---\nEQ/MC01 开关瞬变噪声 测试时间10次\n"
    )
    items = [TestPlanItem(
        code="EQ/MC01", name="开关瞬变噪声", source_location="文件第9页",
        source_quote="模型改写后无法逐字命中的句子",
    )]
    details = [TestPlanDetail(
        code="EQ/MC01", fields={"test_duration": "10次"},
        source_location="文件第9页", source_quote="",
    )]

    repaired = _repair_source_evidence(raw, items, details)

    assert repaired == 2
    assert "EQ/MC01 开关瞬变噪声" in items[0].source_quote
    assert "测试时间10次" in details[0].source_quote
    data = TestPlanData(test_items=items, test_details=details)
    assert validate_extraction(raw, data, require_source_evidence=True) == []


def test_sparse_basic_info_allowed_when_plan_items_exist():
    data = TestPlanData(
        basic_info=TestPlanBasicInfo(part_name="灯具"),
        test_items=[TestPlanItem(code="", name="瞬态抗扰度试验")],
        test_details=[],
    )

    assert enforce_minimum_content(data) == 1


def test_sparse_basic_info_rejected_when_no_plan_content():
    data = TestPlanData(basic_info=TestPlanBasicInfo(part_name="灯具"))

    with pytest.raises(ExtractionError, match="未提取到测试项目"):
        enforce_minimum_content(data)


@pytest.mark.asyncio
async def test_extract_coerces_numeric_llm_fields_for_vendor_xlsx():
    """LLM may return Excel row numbers as ints; extraction should normalize them."""
    from openpyxl import Workbook
    import io

    class FakeReviewer:
        async def _call_api(self, **kwargs):
            return {
                "基本信息": {"零件名称": "灯具", "样品数量": 3},
                "测试项目": [
                    {
                        "编号": 1,
                        "名称": "瞬态抗扰度试验",
                        "是否执行": True,
                        "测试模式": 1,
                        "认可依据": "企业标准",
                        "验收标准": ["必做"],
                    }
                ],
                "测试细则": [
                    {"编号": 1, "工作模式": 1, "测试时间": 10}
                ],
            }

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "验证计划"
    ws["A2"] = "序号"
    ws["B2"] = "试验项目"
    ws["C2"] = "试验描述"
    ws["A3"] = 1
    ws["B3"] = "瞬态抗扰度试验"
    ws["C3"] = "测试灯具瞬态抗扰度"
    buf = io.BytesIO()
    wb.save(buf)

    data = await TestPlanExtractor.extract(
        buf.getvalue(),
        "大冶-灯具准入实验标准（新.xlsx",
        reviewer=FakeReviewer(),
    )

    assert data.basic_info.sample_count == "3"
    assert data.test_items[0].code == "瞬态抗扰度试验"
    assert data.test_items[0].is_executed == ""
    assert data.test_items[0].test_mode == "1"
    assert data.test_items[0].acceptance == '["必做"]'
    assert data.test_details[0].code == "瞬态抗扰度试验"
    assert data.test_details[0].fields["test_duration"] == "10"


def test_execution_flags_are_kept_when_source_has_row_evidence():
    from services.test_plan_extractor import _remove_unsupported_execution_flags

    items = [TestPlanItem(code="A", is_executed="Y"), TestPlanItem(code="B", is_executed="N")]

    removed = _remove_unsupported_execution_flags("A\tY\nB\tN", items)

    assert removed == 0
    assert [item.is_executed for item in items] == ["Y", "N"]


def test_recovers_per_item_sample_counts_from_generic_tabular_columns():
    text = (
        "序号\t试验项目\t技术要求\tRD20容量5000\tRD30容量9000\n"
        "1\t反向电压\t按照标准执行\t3pcs\t2pcs\n"
        "2\t开路测试\t逐路断开\t3pcs\t/"
    )
    items = [TestPlanItem(name="反向电压"), TestPlanItem(name="开路测试")]

    recovered = _recover_tabular_sample_requirements(text, items)

    assert recovered == 2
    assert items[0].sample_requirements == {
        "RD20容量5000": "3pcs", "RD30容量9000": "2pcs",
    }
    assert items[1].sample_requirements == {"RD20容量5000": "3pcs"}


def test_execution_flags_are_removed_when_source_columns_are_blank():
    from services.test_plan_extractor import _remove_unsupported_execution_flags

    items = [TestPlanItem(code="A", is_executed="Y"), TestPlanItem(code="B", is_executed="Y")]

    removed = _remove_unsupported_execution_flags(
        "验证计划\t实际测试情况\nA\t试验描述\nB\t试验描述", items,
    )

    assert removed == 2
    assert [item.is_executed for item in items] == ["", ""]


def test_detail_row_numbers_recovered_only_for_complete_continuous_table():
    items = [
        TestPlanItem(code="耐久试验", name="耐久试验"),
        TestPlanItem(code="防水试验", name="防水试验"),
    ]
    details = [TestPlanDetail(code="1"), TestPlanDetail(code="2")]

    linked = _link_test_details_to_items(details, items)

    assert [detail.code for detail in linked] == ["耐久试验", "防水试验"]


def test_detail_row_numbers_not_guessed_when_counts_do_not_match():
    items = [
        TestPlanItem(code="耐久试验", name="耐久试验"),
        TestPlanItem(code="防水试验", name="防水试验"),
    ]
    details = [TestPlanDetail(code="1")]

    linked = _link_test_details_to_items(details, items)

    assert linked[0].code == "1"


def test_extraction_error():
    err = ExtractionError("msg", [{"f": "x"}])
    assert "msg" in str(err)
    assert len(err.failures) == 1


def test_invalid_format():
    with pytest.raises(ValueError, match="不支持"):
        extract_text(b"x", "test.txt")


def test_models():
    info = TestPlanBasicInfo(part_name="x", part_number="P1")
    assert info.part_name == "x"
    item = TestPlanItem(code="EQ/MC01", is_executed="Y")
    assert item.is_executed == "Y"
    detail = TestPlanDetail(code="EQ/MC01", fields={"k": "v"})
    assert detail.fields["k"] == "v"
    data = TestPlanData(basic_info=info, test_items=[item], test_details=[detail])
    assert len(data.test_items) == 1
