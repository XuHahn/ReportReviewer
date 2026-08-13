import base64
import io
from types import SimpleNamespace
import zipfile

import fitz
from PIL import Image
from openpyxl import Workbook, load_workbook

from services import document_unitizer
from services.document_unitizer import locate_layout_quote, unitize_document


def _pdf(text: str) -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    data = document.tobytes()
    document.close()
    return data


def test_pdf_unitizer_produces_native_and_visual_page():
    result = unitize_document(
        _pdf("Test mode: mode 1"), doc_id="doc-1",
        doc_type="original_records", filename="record.pdf",
    )
    assert result.coverage.state == "complete"
    assert result.coverage.expected_pages == 1
    assert result.coverage.parsed_pages == 1
    assert "Test mode" in result.units[0].native_text
    assert result.units[0].image_data_urls[0].startswith("data:image/png;base64,")
    assert len(result.units[0].source_hash) == 64
    assert len(result.units[0].rendered_pdf_hash) == 64
    assert result.units[0].page_width > 0
    assert result.units[0].page_height > 0
    assert result.units[0].layout_lines[0]["text"] == "Test mode: mode 1"
    assert len(result.units[0].layout_lines[0]["bbox"]) == 4


def test_pdf_unitizer_uses_configured_render_scale(monkeypatch):
    monkeypatch.setenv("UNITIZER_RENDER_SCALE", "2.0")
    result = unitize_document(
        _pdf("render scale"), doc_id="doc-1",
        doc_type="original_records", filename="record.pdf",
    )
    image_bytes = base64.b64decode(
        result.units[0].image_data_urls[0].split(",", 1)[1]
    )
    image = Image.open(io.BytesIO(image_bytes))
    assert image.width > 1000


def test_layout_quote_uses_adjacent_field_label_to_disambiguate_short_value():
    layout = [
        {"text": "Sample name", "bbox": [20, 20, 100, 32]},
        {"text": "AB", "bbox": [120, 20, 140, 32]},
        {"text": "Status AB normal", "bbox": [20, 80, 140, 92]},
    ]

    bbox, anchor_quote = locate_layout_quote(
        layout,
        "AB",
        context_terms=("sample name",),
    )

    assert bbox == [120.0, 20.0, 140.0, 32.0]
    assert anchor_quote == "AB"


def test_layout_quote_keeps_repeated_value_unresolved_without_context():
    layout = [
        {"text": "AB", "bbox": [120, 20, 140, 32]},
        {"text": "Status AB normal", "bbox": [20, 80, 140, 92]},
    ]

    assert locate_layout_quote(layout, "AB") == ([], "")


def test_layout_quote_returns_unresolved_when_only_context_label_matches():
    layout = [
        {"text": "Sample name", "bbox": [20, 20, 100, 32]},
        {"text": "Other value", "bbox": [120, 20, 180, 32]},
    ]

    assert locate_layout_quote(
        layout,
        "missing value",
        context_terms=("sample name",),
    ) == ([], "")


def test_zip_unitizer_preserves_member_and_page_coverage():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("a.pdf", _pdf("record A"))
        archive.writestr("nested/b.pdf", _pdf("record B"))
    result = unitize_document(
        buffer.getvalue(), doc_id="doc-zip", doc_type="original_records",
        filename="records.zip",
    )
    assert result.coverage.state == "complete"
    assert result.coverage.expected_pages == 2
    assert result.coverage.parsed_pages == 2
    assert len(result.units) == 2
    assert all(unit.unit_id.startswith("member-") for unit in result.units)


def test_zip_unitizer_isolates_invalid_member_and_continues():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("valid-a.pdf", _pdf("record A"))
        archive.writestr("broken.pdf", b"%PDF-1.4\nno valid PDF objects")
        archive.writestr("valid-b.pdf", _pdf("record B"))

    result = unitize_document(
        buffer.getvalue(), doc_id="doc-zip", doc_type="original_records",
        filename="records.zip",
    )

    assert result.coverage.state == "partial"
    assert result.coverage.expected_pages == 2
    assert result.coverage.parsed_pages == 2
    assert len(result.units) == 2
    assert result.units[0].unit_id == "member-1/page-1"
    assert result.units[1].unit_id == "member-3/page-1"
    assert result.errors == [
        "压缩包内第 2 个 PDF 文件无法读取（FileDataError）",
    ]


def test_zip_unitizer_ignores_macos_resource_forks():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("record.pdf", _pdf("record"))
        archive.writestr("__MACOSX/._record.pdf", b"AppleDouble metadata")
        archive.writestr("__MACOSX/.DS_Store", b"Finder metadata")

    result = unitize_document(
        buffer.getvalue(), doc_id="doc-zip", doc_type="original_records",
        filename="records.zip",
    )

    assert result.coverage.state == "complete"
    assert len(result.units) == 1
    assert result.errors == []
    assert result.coverage.metadata["ignored_zip_metadata_count"] == 2


def test_zip_unitizer_rejects_path_traversal():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("../escape.pdf", _pdf("bad"))
    try:
        unitize_document(
            buffer.getvalue(), doc_id="doc-zip", doc_type="original_records",
            filename="records.zip",
        )
    except ValueError as exc:
        assert "不安全路径" in str(exc)
    else:
        raise AssertionError("unsafe ZIP path was accepted")


def test_legacy_gbk_zip_filename_is_recovered_for_display():
    expected = "系统12V电源电压波动试验_原始记录.pdf"
    mojibake = expected.encode("gbk").decode("cp437")
    info = zipfile.ZipInfo(mojibake)
    info.flag_bits = 0
    assert document_unitizer.display_zip_filename(info) == expected


def test_utf8_zip_filename_is_never_reinterpreted():
    info = zipfile.ZipInfo("原始记录.pdf")
    info.flag_bits = 0x800
    assert document_unitizer.display_zip_filename(info) == "原始记录.pdf"


def test_utf8_flag_does_not_preserve_obvious_cp437_mojibake():
    expected = "供电电压瞬时下降_原始记录.pdf"
    info = zipfile.ZipInfo(expected.encode("gbk").decode("cp437"))
    info.flag_bits = 0x800
    assert document_unitizer.display_zip_filename(info) == expected


def test_office_units_keep_native_text_page_local(monkeypatch):
    workbook = Workbook()
    workbook.active["A1"] = "WORKBOOK_ONLY_MARKER"
    source = io.BytesIO()
    workbook.save(source)

    rendered = fitz.open()
    first = rendered.new_page()
    first.insert_text((72, 72), "PAGE_ONE")
    second = rendered.new_page()
    second.insert_text((72, 72), "PAGE_TWO")
    rendered.new_page()
    rendered_bytes = rendered.tobytes()
    rendered.close()
    monkeypatch.setattr(document_unitizer, "_office_to_pdf", lambda *_args: rendered_bytes)

    result = unitize_document(
        source.getvalue(), doc_id="doc-plan", doc_type="test_plan",
        filename="plan.xlsx",
    )

    assert [unit.native_text.strip() for unit in result.units] == ["PAGE_ONE", "PAGE_TWO", ""]
    assert all("WORKBOOK_ONLY_MARKER" not in unit.native_text for unit in result.units)


def test_office_conversion_failure_falls_back_to_native_text_without_image(monkeypatch):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["项目", "条件", "结果"])
    sheet.append(["反向电压", "14V", "Pass"])
    source = io.BytesIO()
    workbook.save(source)
    monkeypatch.setattr(
        document_unitizer, "_office_to_pdf",
        lambda *_args: (_ for _ in ()).throw(RuntimeError("LibreOffice unavailable")),
    )

    result = unitize_document(
        source.getvalue(), doc_id="doc-plan", doc_type="test_plan",
        filename="plan.xlsx",
    )

    assert len(result.units) == 1
    assert result.units[0].unit_id == "native-structure"
    assert "反向电压" in result.units[0].native_text
    assert result.units[0].image_data_urls == []
    assert result.errors == ["LibreOffice unavailable"]


def test_xlsx_visual_render_fits_wide_sheets_without_changing_values():
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "test name"
    sheet["Q1"] = "acceptance"
    source = io.BytesIO()
    workbook.save(source)

    rendered_bytes, normalized_sheets = (
        document_unitizer._prepare_xlsx_for_visual_render(source.getvalue())
    )
    rendered = load_workbook(io.BytesIO(rendered_bytes), data_only=False)
    rendered_sheet = rendered.active

    assert normalized_sheets == 1
    assert rendered_sheet["A1"].value == "test name"
    assert rendered_sheet["Q1"].value == "acceptance"
    assert rendered_sheet.sheet_properties.pageSetUpPr.fitToPage is True
    assert rendered_sheet.page_setup.fitToWidth == 1
    assert rendered_sheet.page_setup.fitToHeight == 0
    assert rendered_sheet.page_setup.orientation == "landscape"


def test_office_conversion_passes_fontconfig_and_isolated_profile(monkeypatch):
    captured = {}
    rendered = _pdf("rendered")

    def fake_run(command, **kwargs):
        captured["command"] = command
        captured["env"] = kwargs["env"]
        source = command[-1]
        with open(source.rsplit(".", 1)[0] + ".pdf", "wb") as output:
            output.write(rendered)
        return SimpleNamespace(returncode=0, stderr=b"")

    monkeypatch.setattr(document_unitizer.shutil, "which", lambda _name: "/usr/bin/soffice")
    monkeypatch.setattr(document_unitizer.subprocess, "run", fake_run)
    monkeypatch.setenv("OFFICE_FONTCONFIG", "/opt/fonts/office.conf")
    result = document_unitizer._office_to_pdf(b"xlsx", "plan.xlsx")
    assert result == rendered
    assert captured["env"]["FONTCONFIG_FILE"] == "/opt/fonts/office.conf"
    assert any(item.startswith("-env:UserInstallation=file://") for item in captured["command"])
