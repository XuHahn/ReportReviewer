import io
import json
import zipfile

from openpyxl import load_workbook

from services.evidence_graph_exporter import generate_evidence_package, generate_review_excel
from services.evidence_graph_models import (
    EvidenceRecord,
    FindingStatus,
    GraphNode,
    GraphRun,
    GraphScope,
    GraphSnapshot,
    NodeType,
    ReviewFinding,
)


def _snapshot() -> GraphSnapshot:
    return GraphSnapshot(
        run=GraphRun(graph_id="graph-1", scope=GraphScope.RUN, set_id="set-1", status="machine_complete"),
        evidence=[EvidenceRecord(
            evidence_id="ev-1", graph_id="graph-1", doc_id="report-1",
            doc_type="final_report", filename="=报告.pdf", page_number=8,
            exact_quote="=脉冲2a：符合", extraction_method="native_pdf", confidence=1,
        )],
        nodes=[GraphNode(
            node_id="plan-p2a", graph_id="graph-1", node_type=NodeType.TEST_REQUIREMENT,
            label="电源线瞬态传导抗扰P2a",
        )],
        findings=[ReviewFinding(
            finding_id="finding-1", graph_id="graph-1", check_id="GRAPH-COVERAGE-001",
            status=FindingStatus.CONFIRMED_PASS, title="P2a证据完整",
            subject_node_ids=["plan-p2a"], evidence_ids=["ev-1"], dedupe_key="p2a",
        )],
    )


def test_excel_export_contains_findings_and_neutralizes_formula_prefix():
    content = generate_review_excel(_snapshot())
    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook["审核结论"]

    assert sheet["A2"].value == "1"
    assert sheet["D2"].value == "P2a证据完整"
    assert "=报告.pdf" in sheet["F2"].value
    assert all(
        not str(cell.value or "").startswith("=")
        for cell in sheet[2]
    )


def test_evidence_package_is_self_contained():
    content = generate_evidence_package(_snapshot())
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        assert set(archive.namelist()) == {
            "review-graph.json", "审核结论.xlsx", "审核明细.csv",
        }
        graph_json = archive.read("review-graph.json").decode("utf-8")
        assert "P2a证据完整" in graph_json
        assert '"presentation"' in graph_json


def test_coverage_absence_scope_stays_in_graph_but_not_readable_evidence():
    snapshot = _snapshot()
    snapshot.evidence.append(EvidenceRecord(
        evidence_id="ev-scope", graph_id="graph-1", doc_id="raw-1",
        doc_type="original_records", filename="raw.zip", page_number=1,
        exact_quote="无关检索页", extraction_method="native_pdf", confidence=1,
    ))
    finding = snapshot.findings[0]
    finding.status = FindingStatus.CONFIRMED_ERROR
    finding.title = "P2a未找到原始记录"
    finding.evidence_ids = ["ev-1", "ev-scope"]
    finding.metadata = {"missing_docs": ["original_records"]}

    excel = load_workbook(io.BytesIO(generate_review_excel(snapshot)))
    assert "=脉冲2a：符合" in excel["审核结论"]["H2"].value
    assert "无关检索页" not in excel["审核结论"]["H2"].value

    with zipfile.ZipFile(io.BytesIO(generate_evidence_package(snapshot))) as archive:
        graph = json.loads(archive.read("review-graph.json"))
        assert graph["findings"][0]["evidence_ids"] == ["ev-1", "ev-scope"]
        csv_text = archive.read("审核明细.csv").decode("utf-8-sig")
        assert "无关检索页" not in csv_text
