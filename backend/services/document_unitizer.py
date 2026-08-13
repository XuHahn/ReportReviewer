"""Convert PDF, Office, and ZIP documents into native-text and visual units."""

from __future__ import annotations

import base64
import hashlib
import io
import os
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import zipfile
from pathlib import Path

import fitz
import xlrd
from docx import Document
from openpyxl import load_workbook
from pydantic import BaseModel, Field

from config import MAX_ZIP_FILES, MAX_ZIP_UNCOMPRESSED_BYTES, MAX_ZIP_COMPRESSION_RATIO
from utils.logger import get_logger


logger = get_logger(__name__)


class DocumentUnit(BaseModel):
    unit_id: str
    doc_id: str
    doc_type: str
    filename: str
    native_text: str = ""
    page_number: int = Field(default=0, ge=0)
    sheet_name: str = ""
    cell_range: str = ""
    image_data_urls: list[str] = Field(default_factory=list)
    source_hash: str = ""
    rendered_pdf_hash: str = ""
    page_width: float = Field(default=0.0, ge=0.0)
    page_height: float = Field(default=0.0, ge=0.0)
    page_rotation: int = 0
    layout_lines: list[dict] = Field(default_factory=list)


class CoverageManifest(BaseModel):
    doc_id: str
    doc_type: str
    filename: str = ""
    expected_pages: int = Field(default=0, ge=0)
    parsed_pages: int = Field(default=0, ge=0)
    expected_sheets: int = Field(default=0, ge=0)
    parsed_sheets: int = Field(default=0, ge=0)
    failed_units: list[str] = Field(default_factory=list)
    state: str = "unknown"
    methods: list[str] = Field(default_factory=list)
    metadata: dict = Field(default_factory=dict)


class UnitizationResult(BaseModel):
    coverage: CoverageManifest
    units: list[DocumentUnit] = Field(default_factory=list)
    errors: list[str] = Field(default_factory=list)


def _data_url(png: bytes) -> str:
    return "data:image/png;base64," + base64.b64encode(png).decode("ascii")


def _page_layout_lines(page: fitz.Page) -> list[dict]:
    """Keep compact PDF-point line geometry for later evidence anchoring."""
    lines: list[dict] = []
    for block in page.get_text("dict", sort=True).get("blocks", []):
        for line in block.get("lines", []):
            text = "".join(str(span.get("text") or "") for span in line.get("spans", []))
            bbox = line.get("bbox")
            if not text.strip() or not isinstance(bbox, (list, tuple)) or len(bbox) != 4:
                continue
            lines.append({
                "text": text,
                "bbox": [round(float(value), 3) for value in bbox],
            })
    return lines


def _compact_locator_text(value: str) -> str:
    return "".join(
        character
        for character in unicodedata.normalize("NFKC", str(value or "")).casefold()
        if not character.isspace()
    )


def locate_layout_quote(
    layout_lines: list[dict],
    quote: str,
    *,
    context_terms: tuple[str, ...] = (),
) -> tuple[list[float], str]:
    """Resolve a unique quote to a rendered line bbox without fuzzy matching.

    Context terms are labels such as ``样品名称``. They disambiguate short
    values while the returned anchor quote remains the actual rendered line.
    """
    needle = _compact_locator_text(quote)
    if not needle:
        return [], ""
    contexts = tuple(
        compact for compact in map(_compact_locator_text, context_terms) if compact
    )
    matches: list[tuple[list[float], str, bool]] = []
    for item in layout_lines:
        text = str(item.get("text") or "")
        bbox = item.get("bbox")
        compact = _compact_locator_text(text)
        if needle not in compact or not isinstance(bbox, list) or len(bbox) != 4:
            continue
        matches.append((list(map(float, bbox)), text.strip(), any(term in compact for term in contexts)))
    contextual = [item for item in matches if item[2]]
    if len(contextual) == 1:
        return contextual[0][0], contextual[0][1]
    if len(contextual) > 1:
        return [], ""
    if len(matches) == 1:
        return matches[0][0], matches[0][1]
    if not matches:
        return [], ""
    if not contexts:
        return [], ""

    # PDF tables commonly expose the field label and value as separate text
    # lines even when they are visually in the same row.  Resolve repeated
    # short values against the nearest matching label without fuzzy matching.
    label_boxes = [
        list(map(float, item["bbox"]))
        for item in layout_lines
        if isinstance(item.get("bbox"), list)
        and len(item["bbox"]) == 4
        and any(term in _compact_locator_text(item.get("text") or "") for term in contexts)
    ]
    if not label_boxes:
        return [], ""
    ranked: list[tuple[float, list[float], str]] = []
    for bbox, text, _ in matches:
        center_y = (bbox[1] + bbox[3]) / 2
        score = min(
            abs(center_y - (label[1] + label[3]) / 2)
            + (25.0 if bbox[2] < label[0] else 0.0)
            for label in label_boxes
        )
        ranked.append((score, bbox, text))
    ranked.sort(key=lambda item: item[0])
    height = max(1.0, ranked[0][1][3] - ranked[0][1][1])
    if ranked[0][0] > max(18.0, height * 1.5):
        return [], ""
    if len(ranked) > 1 and abs(ranked[1][0] - ranked[0][0]) < 1.0:
        return [], ""
    return ranked[0][1], ranked[0][2]


def _pdf_units(
    file_bytes: bytes,
    *,
    doc_id: str,
    doc_type: str,
    filename: str,
    unit_prefix: str = "",
    native_fallback: str = "",
    source_hash: str = "",
) -> tuple[list[DocumentUnit], int, list[str]]:
    units: list[DocumentUnit] = []
    errors: list[str] = []
    max_pages = int(os.getenv("UNITIZER_MAX_RENDER_PAGES", "1000"))
    render_scale = float(os.getenv("UNITIZER_RENDER_SCALE", "2.0"))
    if not 1.0 <= render_scale <= 4.0:
        raise ValueError("UNITIZER_RENDER_SCALE 必须在 1.0 到 4.0 之间")
    source_hash = source_hash or hashlib.sha256(file_bytes).hexdigest()
    rendered_pdf_hash = hashlib.sha256(file_bytes).hexdigest()
    with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
        expected_pages = len(pdf)
        for index, page in enumerate(pdf):
            if index >= max_pages:
                errors.append(f"页面数超过渲染上限 {max_pages}")
                break
            try:
                text = page.get_text("text") or ""
                pixmap = page.get_pixmap(
                    matrix=fitz.Matrix(render_scale, render_scale), alpha=False,
                )
                image_url = _data_url(pixmap.tobytes("png"))
                unit_id = f"{unit_prefix}page-{index + 1}" if unit_prefix else f"page-{index + 1}"
                units.append(DocumentUnit(
                    unit_id=unit_id,
                    doc_id=doc_id,
                    doc_type=doc_type,
                    filename=filename,
                    native_text=text if text.strip() else native_fallback,
                    page_number=index + 1,
                    image_data_urls=[image_url],
                    source_hash=source_hash,
                    rendered_pdf_hash=rendered_pdf_hash,
                    page_width=float(page.rect.width),
                    page_height=float(page.rect.height),
                    page_rotation=int(page.rotation),
                    layout_lines=_page_layout_lines(page),
                ))
            except Exception as exc:
                errors.append(f"第 {index + 1} 页渲染失败: {str(exc)[:120]}")
    return units, expected_pages, errors


def _docx_native_text(file_bytes: bytes) -> str:
    document = Document(io.BytesIO(file_bytes))
    chunks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table_index, table in enumerate(document.tables, 1):
        chunks.append(f"## 表格 {table_index}")
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                chunks.append("\t".join(values))
    return "\n".join(chunks)


def _xlsx_native_text(file_bytes: bytes) -> tuple[str, int]:
    workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    chunks: list[str] = []
    for sheet in workbook.worksheets:
        chunks.append(f"## 工作表 {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value).strip() for value in row]
            if any(values):
                chunks.append("\t".join(values))
    return "\n".join(chunks), len(workbook.worksheets)


def _xls_native_text(file_bytes: bytes) -> tuple[str, int]:
    workbook = xlrd.open_workbook(file_contents=file_bytes)
    chunks: list[str] = []
    for sheet in workbook.sheets():
        chunks.append(f"## 工作表 {sheet.name}")
        for row_index in range(sheet.nrows):
            values = [str(sheet.cell_value(row_index, col)).strip() for col in range(sheet.ncols)]
            if any(values):
                chunks.append("\t".join(values))
    return "\n".join(chunks), workbook.nsheets


def _prepare_xlsx_for_visual_render(file_bytes: bytes) -> tuple[bytes, int]:
    """Create a temporary fit-to-width view without changing source content."""
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=False, keep_links=True)
    normalized_sheets = 0
    for sheet in workbook.worksheets:
        value_columns = [
            cell.column
            for row in sheet.iter_rows()
            for cell in row
            if cell.value not in (None, "")
        ]
        if not value_columns:
            continue
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_setup.scale = None
        if max(value_columns) - min(value_columns) + 1 >= 9:
            sheet.page_setup.orientation = "landscape"
        normalized_sheets += 1
    if not normalized_sheets:
        return file_bytes, 0
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue(), normalized_sheets


def _office_to_pdf(file_bytes: bytes, filename: str) -> bytes:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("未安装 LibreOffice，无法生成 Office 视觉页")
    suffix = Path(filename).suffix.lower()
    with tempfile.TemporaryDirectory(prefix="review_office_") as temp_dir:
        source = Path(temp_dir) / (Path(filename).stem + suffix)
        render_bytes = file_bytes
        if suffix == ".xlsx":
            try:
                render_bytes, normalized_sheets = _prepare_xlsx_for_visual_render(
                    file_bytes,
                )
                logger.info(
                    "spreadsheet_render_normalized",
                    normalized_sheet_count=normalized_sheets,
                )
            except Exception as exc:
                logger.warning(
                    "spreadsheet_render_normalization_failed",
                    error_type=type(exc).__name__,
                )
        source.write_bytes(render_bytes)
        profile = Path(temp_dir) / "libreoffice-profile"
        environment = os.environ.copy()
        configured_fontconfig = os.getenv("OFFICE_FONTCONFIG", "").strip()
        bundled_fontconfig = (
            Path(__file__).resolve().parents[2]
            / "scripts/model_runtime/fontconfig-macos-office.conf"
        )
        if configured_fontconfig:
            environment["FONTCONFIG_FILE"] = configured_fontconfig
        elif sys.platform == "darwin" and bundled_fontconfig.is_file():
            environment["FONTCONFIG_FILE"] = str(bundled_fontconfig)
        completed = subprocess.run(
            [
                soffice, f"-env:UserInstallation={profile.as_uri()}",
                "--headless", "--convert-to", "pdf", "--outdir", temp_dir,
                str(source),
            ],
            stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=120,
            check=False, env=environment,
        )
        output = source.with_suffix(".pdf")
        if completed.returncode != 0 or not output.exists():
            stderr = completed.stderr.decode("utf-8", errors="ignore")[:200]
            raise RuntimeError(f"Office 转 PDF 失败: {stderr or completed.returncode}")
        return output.read_bytes()


def _validate_zip(infos: list[zipfile.ZipInfo]) -> None:
    if len(infos) > MAX_ZIP_FILES:
        raise ValueError(f"ZIP 文件数超过上限 {MAX_ZIP_FILES}")
    total = sum(info.file_size for info in infos)
    if total > MAX_ZIP_UNCOMPRESSED_BYTES:
        raise ValueError("ZIP 解压后大小超过上限")
    for info in infos:
        path = Path(info.filename)
        if path.is_absolute() or ".." in path.parts:
            raise ValueError("ZIP 包含不安全路径")
        ratio = info.file_size / max(info.compress_size, 1)
        if ratio > MAX_ZIP_COMPRESSION_RATIO:
            raise ValueError("ZIP 压缩比异常")


def _is_ignored_zip_metadata(info: zipfile.ZipInfo) -> bool:
    """Identify operating-system metadata that is not a business document."""
    path = Path(info.filename)
    return (
        "__MACOSX" in path.parts
        or path.name.startswith("._")
        or path.name in {".DS_Store", "Thumbs.db"}
    )


def _filename_quality(value: str) -> int:
    score = 0
    for char in value:
        codepoint = ord(char)
        category = unicodedata.category(char)
        if 0x4E00 <= codepoint <= 0x9FFF:
            score += 4
        elif char.isascii() and (char.isalnum() or char in " ._-/()[]"):
            score += 1
        elif category.startswith(("L", "N")):
            score += 2
        elif 0x2500 <= codepoint <= 0x259F:
            score -= 5
        elif category.startswith("C"):
            score -= 8
    return score


def display_zip_filename(info: zipfile.ZipInfo) -> str:
    """Recover legacy ZIP metadata without changing the member lookup key.

    Some archives incorrectly mark an already-mojibaked CP437 filename as
    UTF-8.  Keep valid UTF-8 names untouched, but still repair names whose
    character-quality score clearly indicates box-drawing mojibake.
    """
    original = info.filename
    if original.isascii():
        return original
    original_quality = _filename_quality(original)
    if info.flag_bits & 0x800 and original_quality >= 0:
        return original
    try:
        raw = original.encode("cp437")
    except UnicodeEncodeError:
        return original
    encodings = [
        item.strip() for item in os.getenv(
            "ZIP_LEGACY_ENCODINGS", "gb18030,big5,shift_jis",
        ).split(",") if item.strip()
    ]
    best = original
    best_score = original_quality
    for encoding in encodings:
        try:
            candidate = raw.decode(encoding)
        except (LookupError, UnicodeDecodeError):
            continue
        score = _filename_quality(candidate)
        if score > best_score + 2:
            best = candidate
            best_score = score
    return best


# Kept for callers and tests written before the decoder became a shared API.
_display_zip_filename = display_zip_filename


def unitize_document(
    file_bytes: bytes,
    *,
    doc_id: str,
    doc_type: str,
    filename: str,
) -> UnitizationResult:
    lower = filename.lower()
    methods = ["native_structure", "visual_render"]
    errors: list[str] = []
    units: list[DocumentUnit] = []
    expected_pages = parsed_pages = expected_sheets = parsed_sheets = 0
    repaired_zip_filenames = 0
    ignored_zip_metadata = 0

    if lower.endswith(".pdf"):
        units, expected_pages, errors = _pdf_units(
            file_bytes, doc_id=doc_id, doc_type=doc_type, filename=filename,
            source_hash=hashlib.sha256(file_bytes).hexdigest(),
        )
        parsed_pages = len(units)
    elif lower.endswith((".docx", ".xlsx", ".xls")):
        native_text = ""
        if lower.endswith(".docx"):
            native_text = _docx_native_text(file_bytes)
        elif lower.endswith(".xlsx"):
            native_text, expected_sheets = _xlsx_native_text(file_bytes)
            parsed_sheets = expected_sheets
        else:
            native_text, expected_sheets = _xls_native_text(file_bytes)
            parsed_sheets = expected_sheets
        try:
            pdf_bytes = _office_to_pdf(file_bytes, filename)
            units, expected_pages, render_errors = _pdf_units(
                pdf_bytes, doc_id=doc_id, doc_type=doc_type, filename=filename,
                source_hash=hashlib.sha256(file_bytes).hexdigest(),
            )
            parsed_pages = len(units)
            errors.extend(render_errors)
        except Exception as exc:
            errors.append(str(exc)[:300])
            units.append(DocumentUnit(
                unit_id="native-structure", doc_id=doc_id, doc_type=doc_type,
                filename=filename, native_text=native_text,
                source_hash=hashlib.sha256(file_bytes).hexdigest(),
            ))
    elif lower.endswith(".zip"):
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
            infos = [info for info in archive.infolist() if not info.is_dir()]
            _validate_zip(infos)
            ignored_zip_metadata = sum(
                1 for info in infos if _is_ignored_zip_metadata(info)
            )
            supported = [
                info for info in infos
                if not _is_ignored_zip_metadata(info)
                and info.filename.lower().endswith((".pdf", ".docx", ".xlsx", ".xls"))
            ]
            for member_index, info in enumerate(supported, 1):
                display_filename = display_zip_filename(info)
                repaired_zip_filenames += int(display_filename != info.filename)
                try:
                    member_bytes = archive.read(info)
                    member_result = unitize_document(
                        member_bytes, doc_id=doc_id, doc_type=doc_type,
                        filename=display_filename,
                    )
                except Exception as exc:
                    # A single corrupt member must not discard every usable
                    # record in the archive.  Keep a stable, non-content error
                    # marker so the graph becomes machine-incomplete while the
                    # remaining members can still be reviewed.
                    suffix = Path(display_filename).suffix.lower().lstrip(".") or "unknown"
                    errors.append(
                        f"压缩包内第 {member_index} 个 {suffix.upper()} 文件无法读取"
                        f"（{type(exc).__name__}）"
                    )
                    continue
                expected_pages += member_result.coverage.expected_pages
                parsed_pages += member_result.coverage.parsed_pages
                expected_sheets += member_result.coverage.expected_sheets
                parsed_sheets += member_result.coverage.parsed_sheets
                prefix = f"member-{member_index}/"
                for unit in member_result.units:
                    unit.unit_id = prefix + unit.unit_id
                    units.append(unit)
                errors.extend(f"{display_filename}: {error}" for error in member_result.errors)
            if not supported:
                errors.append("ZIP 中没有支持的 PDF/Word/Excel 文件")
            methods.append("safe_zip_inventory")
    else:
        raise ValueError(f"不支持的文档格式: {filename}")

    state = "complete" if units and not errors and all(unit.image_data_urls for unit in units) else "partial" if units else "failed"
    coverage = CoverageManifest(
        doc_id=doc_id, doc_type=doc_type, filename=filename,
        expected_pages=expected_pages, parsed_pages=parsed_pages,
        expected_sheets=expected_sheets, parsed_sheets=parsed_sheets,
        failed_units=errors, state=state, methods=methods,
        metadata={
            "unit_count": len(units),
            "dual_input_complete": state == "complete",
            "repaired_zip_filename_count": repaired_zip_filenames,
            "ignored_zip_metadata_count": ignored_zip_metadata,
        },
    )
    return UnitizationResult(coverage=coverage, units=units, errors=errors)
