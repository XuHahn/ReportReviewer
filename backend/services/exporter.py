import io, os, platform, shutil, subprocess, tempfile

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from constants import SEV_COLORS, SEV_LABELS, RESULT_LABEL, SEV_RGB
from utils.logger import get_logger

logger = get_logger(__name__)


def _find_chrome() -> str:
    """Auto-detect Chrome/Chromium path on macOS, Linux, and Windows."""
    system = platform.system()

    if system == "Darwin":
        candidates = [
            "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
            "/Applications/Chromium.app/Contents/MacOS/Chromium",
        ]
    elif system == "Linux":
        candidates = [
            "google-chrome", "google-chrome-stable", "chromium",
            "chromium-browser", "google-chrome-beta",
        ]
    elif system == "Windows":
        candidates = [
            "C:\\Program Files\\Google\\Chrome\\Application\\chrome.exe",
            "C:\\Program Files (x86)\\Google\\Chrome\\Application\\chrome.exe",
            os.path.expandvars("%LOCALAPPDATA%\\Google\\Chrome\\Application\\chrome.exe"),
        ]
    else:
        candidates = ["google-chrome"]

    for path in candidates:
        if shutil.which(path) or os.path.isfile(path):
            return path
    return candidates[0]

CHROME_PATH = _find_chrome()


def _build_summary_rows(review_items: list[dict]) -> str:
    rows = ""
    for i, item in enumerate(review_items):
        color = SEV_COLORS.get(item["severity"], "#333")
        rows += f"""<tr>
      <td style="text-align:center">{i + 1}</td>
      <td style="color:{color};font-weight:600">{item["severity"]}</td>
      <td>{item["location"]}</td>
      <td>{item["original_text"]}</td>
      <td>{item["error_description"]}</td>
      <td>{item["standard_reference"]}</td>
      <td>{item["suggestion"]}</td></tr>"""
    return rows


def _pdf_html(report: dict) -> str:
    items = report["review_items"]
    filename = report["filename"]
    overall = report["overall_result"]
    result_label = RESULT_LABEL.get(overall, overall)
    rows = _build_summary_rows(items)
    highlighted_html = report["highlighted_html"]

    return f"""<!DOCTYPE html>
<html lang="zh"><head><meta charset="utf-8"><style>
  @page {{ size: A4; margin: 2cm; }}
  body {{ font-family: "PingFang SC", "Microsoft YaHei", sans-serif; font-size: 12px; color: #333; line-height: 1.6; }}
  h1 {{ font-size: 20px; border-bottom: 2px solid #0f3460; padding-bottom: 8px; color: #0f3460; }}
  h2 {{ font-size: 16px; color: #1a1a2e; margin-top: 24px; }}
  .meta {{ color: #667085; font-size: 12px; margin-bottom: 20px; }}
  .badge {{ display: inline-block; padding: 2px 10px; border-radius: 12px; font-weight: 600; font-size: 13px; }}
  .badge-fail {{ background: #fef3f2; color: #b42318; }}
  .badge-pass {{ background: #ecfdf3; color: #067647; }}
  .badge-warning {{ background: #fffaeb; color: #b54708; }}
  .badge-error {{ background: #fef3f2; color: #b42318; }}
  table.summary {{ width: 100%; border-collapse: collapse; font-size: 10px; margin: 16px 0; }}
  table.summary th {{ background: #0f3460; color: #fff; padding: 8px 6px; text-align: left; font-weight: 500; }}
  table.summary td {{ padding: 6px; border: 1px solid #e4e7ec; vertical-align: top; }}
  table.summary tr:nth-child(even) {{ background: #fafbfc; }}
  .error-highlight {{ padding: 1px 2px; border-radius: 2px; font-weight: 500; }}
  .error-error {{ background-color: #fecdca; border-bottom: 2px solid #d92d20; color: #b42318; }}
  .error-warning {{ background-color: #fde272; border-bottom: 2px solid #dc6803; color: #b54708; }}
  .error-info {{ background-color: #b2ddff; border-bottom: 2px solid #2e90fa; color: #175cd3; }}
  .error-idx-badge {{ display: inline-block; color: #fff; font-size: 8px; font-weight: 700;
    width: 14px; height: 14px; line-height: 14px; text-align: center; border-radius: 50%; margin-left: 1px; vertical-align: super; }}
  .error-highlight.error-error .error-idx-badge {{ background: #d92d20; }}
  .error-highlight.error-warning .error-idx-badge {{ background: #dc6803; }}
  .error-highlight.error-info .error-idx-badge {{ background: #2e90fa; }}
  .report-preview h1 {{ font-size: 18px; margin: 16px 0 8px; }}
  .report-preview h2 {{ font-size: 15px; margin: 14px 0 6px; }}
  .report-preview h3 {{ font-size: 13px; margin: 10px 0 4px; }}
  .report-preview table {{ border-collapse: collapse; width: 100%; font-size: 10px; margin: 8px 0; }}
  .report-preview td, .report-preview th {{ border: 1px solid #e4e7ec; padding: 4px 6px; }}
  .report-preview th {{ background: #f9fafb; font-weight: 600; }}
  .report-preview img {{ max-width: 100%; }}
  .page-break {{ page-break-before: always; }}
</style></head><body>
<h1>EMC检测报告审核结果</h1>
<p class="meta">
  文件：{filename}<br>
  审核结论：<span class="badge badge-{overall}">{result_label}</span><br>
  问题总数：{len(items)} 处<br>
  审核时间：{report["created_at"][:19]}
</p>
<h2>审核问题汇总表</h2>
<table class="summary"><thead><tr>
  <th>#</th><th>级别</th><th>位置</th><th>原文</th><th>问题描述</th><th>标准依据</th><th>建议</th>
</tr></thead><tbody>{rows}</tbody></table>
<div class="page-break"></div>
<h2>报告原文（含标红标注）</h2>
<div class="report-preview">{highlighted_html}</div>
</body></html>"""


def generate_pdf(report: dict) -> bytes:
    """Generate PDF from report record. Returns PDF bytes."""
    html = _pdf_html(report)
    with tempfile.NamedTemporaryFile(suffix=".html", mode="w", encoding="utf-8", delete=False) as f:
        f.write(html)
        html_path = f.name

    pdf_path = html_path + ".pdf"
    try:
        base_args = [CHROME_PATH, "--headless", "--disable-gpu",
                     f"--print-to-pdf={pdf_path}", f"file://{html_path}"]
        try:
            subprocess.run(base_args, check=True, timeout=30, capture_output=True)
        except subprocess.CalledProcessError:
            logger.warning("Chrome w/o --no-sandbox failed, retrying with --no-sandbox")
            subprocess.run([CHROME_PATH, "--headless", "--disable-gpu", "--no-sandbox",
                            f"--print-to-pdf={pdf_path}", f"file://{html_path}"],
                           check=True, timeout=30, capture_output=True)

        with open(pdf_path, "rb") as f:
            return f.read()
    finally:
        if os.path.exists(html_path):
            os.unlink(html_path)
        if os.path.exists(pdf_path):
            os.unlink(pdf_path)


# ── Word export ───────────────────────────────────────────────────────

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _shade_cell(cell, color_hex: str):
    s = OxmlElement("w:shd")
    s.set(qn("w:fill"), color_hex)
    s.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(s)


def _add_run(para, text, *, bold=False, size=Pt(10.5), color=None):
    r = para.add_run(text)
    r.bold = bold
    r.font.size = size
    if color:
        r.font.color.rgb = color
    return r


def generate_docx(report: dict) -> bytes:
    """Generate Word docx (error list only, no original report)."""
    items = report["review_items"]
    filename = report["filename"]
    overall = report["overall_result"]
    result_label = RESULT_LABEL.get(overall, overall)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    style = doc.styles["Normal"]
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.35

    # cover
    title = doc.add_heading("EMC检测报告审核结果", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    result_color = {"pass": RGBColor(6, 118, 71), "fail": RGBColor(180, 35, 24),
                    "warning": RGBColor(181, 71, 8)}.get(overall, RGBColor(0, 0, 0))
    for label, value in [
        ("文件名称", filename),
        ("审核结论", result_label),
        ("问题总数", f"{len(items)} 处"),
        ("审核时间", report["created_at"][:19]),
    ]:
        p = doc.add_paragraph()
        _add_run(p, f"{label}：", bold=True, size=Pt(11))
        if label == "审核结论":
            _add_run(p, value, bold=True, size=Pt(11), color=result_color)
        else:
            _add_run(p, value, size=Pt(11))

    counts = {"error": 0, "warning": 0, "info": 0}
    for item in items:
        counts[item["severity"]] = counts.get(item["severity"], 0) + 1
    p = doc.add_paragraph()
    _add_run(p, "严重程度分布：", bold=True, size=Pt(11))
    for key in ["error", "warning", "info"]:
        _add_run(p, f"  {SEV_LABELS[key]} {counts[key]}  ", size=Pt(11), color=SEV_RGB[key])
    doc.add_paragraph()

    # summary table
    doc.add_heading("一、审核问题汇总", level=2)
    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, text in enumerate(["#", "级别", "所在位置", "原文", "问题描述", "标准依据", "修改建议"]):
        table.rows[0].cells[j].text = ""
        p = table.rows[0].cells[j].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_run(p, text, bold=True, size=Pt(9), color=RGBColor(255, 255, 255))
        _shade_cell(table.rows[0].cells[j], "0F3460")

    for i, item in enumerate(items):
        row = table.add_row()
        vals = [str(i + 1), item["severity"], item["location"], item["original_text"],
                item["error_description"], item["standard_reference"], item["suggestion"]]
        for j, val in enumerate(vals):
            row.cells[j].text = ""
            p = row.cells[j].paragraphs[0]
            color = SEV_RGB.get(item["severity"]) if j == 1 else None
            _add_run(p, val, bold=(j == 1), size=Pt(8), color=color)

    widths = [Cm(0.7), Cm(1.0), Cm(2.2), Cm(3.0), Cm(3.8), Cm(2.8), Cm(3.8)]
    for row in table.rows:
        for j, w in enumerate(widths):
            row.cells[j].width = w
    doc.add_paragraph()

    # per-item detail
    doc.add_heading("二、审核问题逐项详情", level=2)
    for i, item in enumerate(items):
        sev = item["severity"]
        color = SEV_RGB.get(sev, RGBColor(0, 0, 0))
        h = doc.add_heading(level=3)
        _add_run(h, f"#{i + 1}  ", bold=True, size=Pt(12), color=color)
        _add_run(h, f"[{SEV_LABELS.get(sev, sev)}]  ", bold=True, size=Pt(12), color=color)
        _add_run(h, item["location"], size=Pt(12))

        for label, value in [
            ("原文", item["original_text"]),
            ("问题描述", item["error_description"]),
            ("标准依据", item["standard_reference"]),
            ("修改建议", item["suggestion"]),
        ]:
            p = doc.add_paragraph()
            _add_run(p, f"{label}：", bold=True)
            _add_run(p, value)
        doc.add_paragraph()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── Batch Excel Export ─────────────────────────────────────────────────────

HEADER_FILL = PatternFill(start_color="1A1A2E", end_color="1A1A2E", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
RESULT_COLORS = {"pass": "027A48", "fail": "B42318", "warning": "B54708"}
SEV_COLORS_HEX = {"error": "B42318", "warning": "B54708", "info": "175CD3"}
THIN_BORDER = Border(
    left=Side(style="thin", color="D0D5DD"),
    right=Side(style="thin", color="D0D5DD"),
    top=Side(style="thin", color="D0D5DD"),
    bottom=Side(style="thin", color="D0D5DD"),
)


def _write_row(ws, row: int, values: list, bold: bool = False):
    for col, val in enumerate(values, 1):
        cell = ws.cell(row=row, column=col, value=val)
        cell.border = THIN_BORDER
        cell.alignment = Alignment(vertical="center")
        if bold:
            cell.font = Font(bold=True, size=11)
        if row == 1:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT


def generate_batch_excel(reports: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "审核汇总"

    # Summary sheet
    _write_row(ws, 1, ["文件名", "审核结果", "问题数", "上传者", "审核时间"], bold=True)
    for i, r in enumerate(reports):
        row = i + 2
        _write_row(ws, row, [
            r.get("filename", ""),
            RESULT_LABEL.get(r.get("overall_result", ""), r.get("overall_result", "")),
            len(r.get("review_items", [])),
            r.get("employee_id", ""),
            r.get("created_at", ""),
        ])
        result_color = RESULT_COLORS.get(r.get("overall_result", ""))
        if result_color:
            ws.cell(row=row, column=2).font = Font(color=result_color, bold=True)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 20

    # Per-report sheets
    for r in reports:
        rid = r.get("report_id", r.get("id", ""))
        sheet_name = rid[:12] if rid else "unknown"
        ws = wb.create_sheet(title=sheet_name)
        headers = ["严重程度", "位置", "原文", "错误描述", "标准依据", "修改建议"]
        columns = ["severity", "location", "original_text", "error_description",
                   "standard_reference", "suggestion"]
        _write_row(ws, 1, headers, bold=True)
        items = r.get("review_items", [])
        for i, item in enumerate(items):
            row = i + 2
            vals = [item.get(k, "") for k in columns]
            sev = item.get("severity", "")
            label = SEV_LABELS.get(sev, sev)
            vals[0] = label
            _write_row(ws, row, vals)
            sev_color = SEV_COLORS_HEX.get(sev)
            if sev_color:
                ws.cell(row=row, column=1).font = Font(color=sev_color, bold=True)
        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 20
        ws.column_dimensions["C"].width = 45
        ws.column_dimensions["D"].width = 45
        ws.column_dimensions["E"].width = 30
        ws.column_dimensions["F"].width = 40
        ws.sheet_properties.tabColor = "1A1A2E"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
