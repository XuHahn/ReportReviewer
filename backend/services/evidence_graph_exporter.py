"""Human-readable and machine-auditable exports for one evidence graph run."""

from __future__ import annotations

import csv
import html
import io
import json
import os
import re
import subprocess
import tempfile
import zipfile
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from services.evidence_graph_models import GraphSnapshot
from services.finding_presentation import build_finding_presentation
from utils.browser_runtime import CHROME_PATH


# ── 中文标签映射 ──────────────────────────────────────────────────────────

FINDING_STATUS_LABELS = {
    "confirmed_error": "确定问题",
    "unresolved": "待确认",
    "confirmed_advisory": "提醒",
    "unresolved_advisory": "待确认提醒",
    "confirmed_pass": "证据完整",
    "not_applicable": "不适用",
}

SEVERITY_LABELS = {
    "error": "需处理",
    "warning": "需确认",
    "info": "提示",
}

# 严重程度对应的单元格底色
SEVERITY_FILLS = {
    "需处理": PatternFill("solid", fgColor="FDE8E8"),   # 浅红
    "需确认": PatternFill("solid", fgColor="FEF3C7"),   # 浅橙
    "提示": PatternFill("solid", fgColor="DBEAFE"),     # 浅蓝
}

# check_id → 问题类别（中文）
CHECK_CATEGORY_LABELS: dict[str, str] = {
    "DOC-STRUCTURE-001": "文档结构",
    "DOC-STRUCTURE-002": "文档结构",
    "DOC-TIMELINE-001": "时间逻辑",
    "DOC-TIMELINE-002": "时间逻辑",
    "DOC-TIMELINE-003": "时间逻辑",
    "DOC-REQUIRED-FIELD-001": "必填字段",
    "DOC-SIGNATURE-001": "签名审批",
    "DOC-PAGINATION-001": "页码完整性",
    "DOC-REFERENCE-001": "引用完整性",
    "DOC-CROSS-FIELD-001": "跨文档校验",
    "DOC-PLAN-REF-001": "计划一致性",
    "DOC-RESULT-CONSISTENCY-001": "跨文档校验",
    "DOC-INSTRUMENT-CONSISTENCY-001": "仪器设备",
    "RESULT-DIMENSION-001": "结果完整性",
    "RESULT-EMPTY-001": "结果完整性",
    "RESULT-TYPE-001": "数据类型",
    "RESULT-DUPLICATE-001": "数据类型",
    "RESULT-FORMULA-001": "结果计算",
    "RESULT-LIMIT-001": "结果计算",
    "INSTRUMENT-IDENTITY-001": "仪器设备",
    "INSTRUMENT-CAL-001": "仪器设备",
    "ANOMALY-CONCLUSION-001": "异常与结论",
    "EVIDENCE-PROFILE-001": "证据完整性",
    "SEMANTIC-VARIABLE-001": "语义一致性",
    "GRAPH-CONCLUSION-001": "图表校验",
    "GRAPH-COVERAGE-001": "测试项覆盖",
    "GRAPH-COVERAGE-002": "计划外项目",
    "GRAPH-SAMPLE-001": "样品信息",
    "REPORT-RESULT-CONFLICT-001": "报告内校验",
    "REPORT-SPEC-RESULT-001": "报告内校验",
    "PLAN-ACCEPTANCE-CONFLICT-001": "计划一致性",
    "PLAN-CLAUSE-CONSISTENCY-001": "计划一致性",
    "STANDARD-PARAMETER-CONSISTENCY-001": "标准符合性",
    "STANDARD-ACCEPTANCE-CONSISTENCY-001": "标准符合性",
    "STANDARD-PLAN-COVERAGE-001": "标准覆盖",
    "STANDARD-SCOPE-ADVISORY-001": "标准符合性",
}

DOC_TYPE_LABELS: dict[str, str] = {
    "order_form": "委托单",
    "test_plan": "试验计划",
    "original_records": "原始记录",
    "final_report": "检测报告",
    "test_standard": "测试标准",
}

DECISION_LABELS: dict[str, str] = {
    "confirmed": "已确认",
    "dismissed": "已排除",
    "advisory": "已记录",
    "unresolved": "待处理",
}


# ── 辅助函数 ──────────────────────────────────────────────────────────────

def _safe_cell(value: Any) -> str:
    text = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False, default=str)
    text = "".join(char for char in text if char in "\n\t" or ord(char) >= 32)
    if text.startswith(("=", "+", "-", "@")):
        text = "'" + text
    return text[:32000]


def _trim(value: Any) -> str:
    """Safely cast to str and strip whitespace."""
    return str(value).strip() if value else ""


def _format_location(evidence_item) -> str:
    """Format page/sheet location for a single evidence record."""
    parts: list[str] = []
    if evidence_item.page_number:
        parts.append(f"第{evidence_item.page_number}页")
    if evidence_item.sheet_name:
        parts.append(evidence_item.sheet_name)
    if evidence_item.cell_range:
        parts.append(evidence_item.cell_range)
    return " ".join(parts) if parts else ""


def _filename_or_empty(evidence_item) -> str:
    return _trim(getattr(evidence_item, "filename", ""))


def _display_filename(evidence_item) -> str:
    """Resolve a human-readable filename for one evidence record.

    When the stored filename is a .zip archive (common for original_records),
    the real source is a file inside the archive.  We extract the member index
    from metadata.unit_id and show "归档包名 / 内部文件 #N".
    """
    raw = _filename_or_empty(evidence_item)
    if not raw.lower().endswith(".zip"):
        return raw
    # metadata.unit_id looks like "original_records:member-7/page-1"
    unit_id = _trim(getattr(evidence_item, "metadata", {}).get("unit_id", ""))
    match = re.search(r"member-(\d+)", unit_id)
    if match:
        return f"{raw} / 内部文件 #{match.group(1)}"
    return raw


def _doc_label_with_file(evidence_item) -> str:
    """e.g. '检测报告（RE_Report.pdf）' or '原始记录（archive.zip / 内部文件 #3）'."""
    label = DOC_TYPE_LABELS.get(getattr(evidence_item, "doc_type", ""), evidence_item.doc_type)
    display = _display_filename(evidence_item)
    return f"{label}（{display}）" if display else label


def _finding_presentation(finding) -> dict[str, Any]:
    """Use the immutable presentation saved with the finding.

    Older snapshots may not contain it; generate the same deterministic
    fallback at export time so historical exports keep the structured model.
    """
    metadata = getattr(finding, "metadata", {}) or {}
    stored = metadata.get("presentation") if isinstance(metadata, dict) else None
    return stored if isinstance(stored, dict) else build_finding_presentation(finding)


def _llm_supplement_text(presentation: dict[str, Any]) -> str:
    supplement = presentation.get("llm_supplement")
    if not isinstance(supplement, dict) or supplement.get("status") != "ready":
        return ""
    parts = []
    for label, key in (("判断", "judgment"), ("依据", "basis"), ("不确定点", "uncertainty"), ("系统处理", "handling")):
        value = _trim(supplement.get(key))
        if value:
            parts.append(f"{label}：{value}")
    return "\n".join(parts)


def _snapshot_dump_with_presentations(snapshot: GraphSnapshot) -> dict[str, Any]:
    """Export a complete structured snapshot without rewriting stored history."""
    dumped = snapshot.model_dump(mode="json")
    finding_by_id = {finding.finding_id: finding for finding in snapshot.findings}
    for item in dumped.get("findings", []):
        finding = finding_by_id.get(str(item.get("finding_id")))
        if finding is None:
            continue
        metadata = item.get("metadata")
        if not isinstance(metadata, dict):
            metadata = {}
            item["metadata"] = metadata
        if not isinstance(metadata.get("presentation"), dict):
            metadata["presentation"] = _finding_presentation(finding)
    return dumped


def _readable_evidence_items(finding, evidence_map: dict[str, Any]) -> list[Any]:
    """Return direct evidence for human-readable exports.

    Coverage-missing findings retain every searched page as machine-auditable
    absence scope.  Those pages stay in review-graph.json, while the readable
    views mirror V2 and show only evidence that directly identifies the test
    item.
    """
    items = [evidence_map[eid] for eid in finding.evidence_ids if eid in evidence_map]
    if str(finding.check_id) != "GRAPH-COVERAGE-001":
        return items
    metadata = getattr(finding, "metadata", {}) or {}
    missing_docs = {
        str(doc_type) for doc_type in metadata.get("missing_docs", [])
        if isinstance(doc_type, str)
    }
    return [item for item in items if item.doc_type not in missing_docs]


def _plain_title(finding, evidence_items: list) -> str:
    """Return the finding title as-is — the engine already produces clear Chinese.

    The only override is for DOC-CROSS-FIELD-001 where the title ends with
    a generic suffix like "…存在不同写法" and we prepend doc-type context.
    """
    check_id = str(finding.check_id)
    raw_title = _trim(getattr(finding, "title", ""))

    if check_id == "DOC-CROSS-FIELD-001":
        subject = raw_title.split("存在不同写法")[0].strip()
        doc_labels = sorted({_trim(DOC_TYPE_LABELS.get(e.doc_type, e.doc_type)) for e in evidence_items})
        if len(doc_labels) > 1 and subject:
            return f"{'与'.join(doc_labels)}中的{subject}写法不同"
        if subject:
            return f"{subject}在不同资料中的写法不同"

    return raw_title if raw_title else str(getattr(finding, "description", ""))


def _plain_summary(finding, evidence_items: list) -> str:
    """Build a readable summary: engine description first, enriched with evidence
    values for check types where a direct comparison is useful."""
    check_id = str(finding.check_id)
    metadata = getattr(finding, "metadata", {}) or {}
    engine_desc = _trim(getattr(finding, "description", ""))

    # ── Special cases where evidence data adds useful context ──────────

    if check_id == "GRAPH-COVERAGE-001":
        missing = metadata.get("missing_docs", []) if isinstance(metadata.get("missing_docs"), list) else []
        if missing:
            label_map = {"original_records": "原始记录", "final_report": "检测报告", "test_plan": "试验计划"}
            detail = "；".join(label_map.get(d, str(d)) + "未找到对应项" for d in missing)
            return f"{engine_desc}\n补充说明：{detail}" if engine_desc else detail
        # When coverage is complete, the engine description is already clear
        return engine_desc

    if check_id == "GRAPH-COVERAGE-002":
        if metadata.get("ambiguous_identity"):
            detail = "原始记录与检测报告之间的项目身份尚未唯一闭合"
        elif metadata.get("plan_extraction_complete") is False:
            detail = "测试计划范围提取尚未完整，当前只能保留为待确认"
        elif metadata.get("raw_item_id") and metadata.get("report_item_ids"):
            detail = "原始记录和检测报告均出现该项目，但完整测试计划中未找到对应要求"
        elif metadata.get("raw_item_id"):
            detail = "原始记录出现该项目，但完整测试计划中未找到对应要求"
        else:
            detail = "检测报告出现该项目，但完整测试计划中未找到对应要求"
        return f"{engine_desc}\n补充说明：{detail}" if engine_desc else detail

    if check_id == "DOC-CROSS-FIELD-001":
        parts: list[str] = []
        for e in evidence_items:
            label = DOC_TYPE_LABELS.get(getattr(e, "doc_type", ""), e.doc_type)
            value = _trim(getattr(e, "exact_quote", "")).replace("\n", " ").replace("\r", " ")
            compact = (value[:40] + "…") if len(value) > 40 else value
            parts.append(f"{label}填写为「{compact or '未提取到原值'}」")
        comparison = "；".join(dict.fromkeys(parts))
        if engine_desc:
            return f"{engine_desc}\n对比：{comparison}"
        return comparison

    if check_id == "REPORT-SPEC-RESULT-001":
        declared = _trim(metadata.get("declared_value", ""))
        observed = [v for v in metadata.get("observed_values", []) if isinstance(v, str) and v.strip()]
        if declared and observed:
            detail = (
                f"要求等级（Performance criteria）：{declared}；"
                f"实际等级（Actual performance）：{'、'.join(observed)}"
            ) if metadata.get("comparison_kind") == "performance_criteria_vs_actual_performance" else (
                f"规范要求：{declared}；结果记录实际采用：{'、'.join(observed)}"
            )
            return f"{engine_desc}\n{detail}" if engine_desc else detail

    if check_id == "RESULT-EMPTY-001":
        return (
            "报告该试验的汇总结论为通过，但明细结果表格中的判定数据系统未能完整读取。\n"
            "请人工打开检测报告，核对明细部分的判定结果是否已填写。\n"
            "如明细已填写完整 → 可排除此问题（系统提取遗漏）；\n"
            "如明细确实空白 → 需联系试验工程师补充。"
        )

    # ── Default: use the engine's own description (already clear Chinese) ──
    return engine_desc if engine_desc else str(getattr(finding, "title", ""))


def _build_evidence_quote(finding, evidence_items: list) -> str:
    """Build a readable evidence quote block grouped by document."""
    lines: list[str] = []
    for e in evidence_items:
        label = DOC_TYPE_LABELS.get(getattr(e, "doc_type", ""), e.doc_type)
        display = _display_filename(e)
        header = f"【{label}】{display}"
        loc = _format_location(e)
        if loc:
            header += f"  {loc}"
        quote = _trim(getattr(e, "exact_quote", ""))
        if quote:
            lines.append(f"{header}\n  → {quote}")
        else:
            lines.append(header)
    return "\n\n".join(lines) if lines else ""


# ── Excel 导出 ────────────────────────────────────────────────────────────

def generate_review_excel(snapshot: GraphSnapshot) -> bytes:
    workbook = Workbook()

    # ── Sheet 1: 审核结论 ─────────────────────────────────────────────
    sheet = workbook.active
    sheet.title = "审核结论"
    headers = [
        "序号", "问题类别", "严重程度", "问题标题", "问题说明",
        "涉及文档（文件名）", "原文位置", "原文摘录",
        "处理状态", "处理备注", "处理人", "处理时间",
        "系统状态", "核对对象", "已核对范围", "结构化发现", "影响", "LLM补充说明",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="184D76")

    evidence_map = {item.evidence_id: item for item in snapshot.evidence}
    decisions_map = {item.finding_id: item for item in snapshot.decisions}

    for index, finding in enumerate(snapshot.findings, 1):
        evidence_items = _readable_evidence_items(finding, evidence_map)

        category = CHECK_CATEGORY_LABELS.get(str(finding.check_id), str(finding.check_id))
        severity_label = SEVERITY_LABELS.get(str(finding.severity), str(finding.severity))
        title = _plain_title(finding, evidence_items)
        summary = _plain_summary(finding, evidence_items)
        presentation = _finding_presentation(finding)

        # 涉及文档：文档类型（文件名），去重换行
        doc_labels = list(dict.fromkeys(
            _doc_label_with_file(e) for e in evidence_items
        ))
        doc_str = "\n".join(doc_labels) if doc_labels else ""

        # 原文位置：去重换行
        locations = list(dict.fromkeys(
            _format_location(e) for e in evidence_items if _format_location(e)
        ))
        location_str = "\n".join(locations) if locations else ""

        # 原文摘录
        quote_str = _build_evidence_quote(finding, evidence_items)

        # 处理决策
        decision = decisions_map.get(str(finding.finding_id))
        if decision:
            status_str = DECISION_LABELS.get(str(decision.decision), str(decision.decision))
            comment = _trim(getattr(decision, "comment", ""))
            actor = _trim(getattr(decision, "actor_id", ""))
            processed_at = _trim(getattr(decision, "created_at", ""))
        else:
            status_str = "待处理"
            comment = ""
            actor = ""
            processed_at = ""

        sheet.append([
            index,
            category,
            severity_label,
            title,
            summary,
            doc_str,
            location_str,
            quote_str,
            status_str,
            comment,
            actor,
            processed_at,
            presentation.get("status_label", ""),
            presentation.get("subject", ""),
            presentation.get("checked", ""),
            presentation.get("issue", ""),
            presentation.get("impact", ""),
            _llm_supplement_text(presentation),
        ])

    # Column widths
    widths = [6, 12, 10, 42, 50, 30, 16, 60, 10, 40, 12, 18, 14, 28, 34, 50, 42, 52]
    for idx, width in enumerate(widths, 1):
        col_letter = chr(64 + idx) if idx <= 26 else chr(64 + (idx - 1) // 26) + chr(65 + (idx - 1) % 26)
        sheet.column_dimensions[col_letter].width = width

    # Apply styles: severity color fills + wrap text for all cells
    for row in sheet.iter_rows(min_row=2):
        severity_cell = row[2]  # 严重程度 is column 3 (0-indexed: 2)
        severity_text = str(severity_cell.value) if severity_cell.value else ""
        fill = SEVERITY_FILLS.get(severity_text)
        if fill:
            severity_cell.fill = fill
        for cell in row:
            cell.value = _safe_cell(cell.value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # ── Sheet 2: 证据关系 ─────────────────────────────────────────────
    nodes = {item.node_id: item for item in snapshot.nodes}
    relation_sheet = workbook.create_sheet("证据关系")
    relation_sheet.append(["来源节点", "关系", "目标节点", "来源", "置信度", "依据"])
    for cell in relation_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="345B7A")
    for edge in snapshot.edges:
        relation_sheet.append([
            nodes.get(edge.source_node_id).label if edge.source_node_id in nodes else edge.source_node_id,
            str(edge.relation_type),
            nodes.get(edge.target_node_id).label if edge.target_node_id in nodes else edge.target_node_id,
            str(edge.origin), edge.confidence, edge.rationale,
        ])
    for column, width in zip("ABCDEF", [35, 20, 35, 20, 12, 60]):
        relation_sheet.column_dimensions[column].width = width

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _review_html(snapshot: GraphSnapshot) -> str:
    evidence_map = {item.evidence_id: item for item in snapshot.evidence}
    decisions_map = {item.finding_id: item for item in snapshot.decisions}

    # Summary counters
    error_count = sum(1 for f in snapshot.findings if str(f.severity) == "error")
    warning_count = sum(1 for f in snapshot.findings if str(f.severity) == "warning")
    info_count = sum(1 for f in snapshot.findings if str(f.severity) == "info")

    rows: list[str] = []
    for idx, finding in enumerate(snapshot.findings, 1):
        evidence_items = _readable_evidence_items(finding, evidence_map)

        category = CHECK_CATEGORY_LABELS.get(str(finding.check_id), str(finding.check_id))
        severity_label = SEVERITY_LABELS.get(str(finding.severity), str(finding.severity))
        title = html.escape(_plain_title(finding, evidence_items))
        summary = html.escape(_plain_summary(finding, evidence_items))
        presentation = _finding_presentation(finding)
        presentation_rows = "".join(
            f"<div><b>{html.escape(label)}</b><span>{html.escape(_trim(presentation.get(key)))}</span></div>"
            for label, key in (("状态", "status_label"), ("核对对象", "subject"), ("已核对", "checked"), ("发现问题", "issue"), ("影响", "impact"))
            if _trim(presentation.get(key))
        )
        llm_text = html.escape(_llm_supplement_text(presentation))
        supplement_html = f"<section class='supplement'><b>补充判断</b><p>{llm_text}</p></section>" if llm_text else ""

        # Evidence quotes
        quote_parts: list[str] = []
        for e in evidence_items:
            label = DOC_TYPE_LABELS.get(getattr(e, "doc_type", ""), e.doc_type)
            filename = html.escape(_display_filename(e))
            loc = _format_location(e)
            header = f"{label}  {filename}"
            if loc:
                header += f"  {loc}"
            quote = html.escape(_trim(getattr(e, "exact_quote", "")))
            if quote:
                quote_parts.append(f"<blockquote><small>{header}</small>{quote}</blockquote>")
            else:
                quote_parts.append(f"<blockquote><small>{header}</small></blockquote>")
        quotes_html = "".join(quote_parts)

        # Decision
        decision = decisions_map.get(str(finding.finding_id))
        if decision:
            decision_label = DECISION_LABELS.get(str(getattr(decision, "decision", "")), "")
            decision_comment = html.escape(_trim(getattr(decision, "comment", "")))
            decision_actor = html.escape(_trim(getattr(decision, "actor_id", "")))
            decision_time = html.escape(_trim(getattr(decision, "created_at", "")))
            decision_html = (
                f"<footer class='decision'><b>处理结论：{decision_label}</b>"
                f"<p>{decision_comment}</p>"
                f"<small>{decision_actor}  {decision_time}</small></footer>"
            )
        else:
            decision_html = "<footer class='decision pending'><b>待处理</b></footer>"

        css_class = html.escape(str(finding.status))
        rows.append(
            f"<article class='{css_class} severity-{html.escape(str(finding.severity))}'>"
            f"<header><span class='index'>#{idx}</span>"
            f"<span class='category'>{html.escape(category)}</span>"
            f"<span class='severity {html.escape(str(finding.severity))}'>{html.escape(severity_label)}</span></header>"
            f"<h2>{title}</h2><p class='summary'>{summary}</p>"
            f"<section class='presentation'>{presentation_rows}</section>{supplement_html}"
            f"{quotes_html}{decision_html}</article>"
        )

    return f"""<!doctype html><html lang='zh-CN'><meta charset='utf-8'><style>
    @page {{ size:A4; margin:16mm; }}
    body {{ font-family:'Noto Sans CJK SC','PingFang SC','Microsoft YaHei',sans-serif;color:#1e293b;font-size:11px;line-height:1.6 }}
    h1 {{ color:#0f3b5e;border-bottom:2px solid #0f3b5e;padding-bottom:10px;font-size:20px;margin-bottom:6px }}
    .meta {{ color:#64748b;font-size:10px;margin-bottom:24px }}
    .stats {{ display:flex;gap:16px;margin-bottom:24px }}
    .stats div {{ padding:8px 16px;border-radius:6px;font-weight:bold;font-size:13px }}
    .stats .err {{ background:#fee2e2;color:#b91c1c }} .stats .warn {{ background:#fef3c7;color:#92400e }} .stats .info {{ background:#dbeafe;color:#1e40af }}

    article {{ break-inside:avoid;border:1px solid #e2e8f0;border-left:5px solid #94a3b8;border-radius:8px;padding:14px;margin:12px 0 }}
    article.severity-error {{ border-left-color:#dc2626 }}
    article.severity-warning {{ border-left-color:#f59e0b }}
    article.severity-info {{ border-left-color:#3b82f6 }}

    article header {{ display:flex;align-items:center;gap:10px;margin-bottom:6px }}
    .index {{ color:#94a3b8;font-size:10px;min-width:24px }}
    .category {{ background:#f1f5f9;color:#475569;padding:2px 8px;border-radius:4px;font-size:10px }}
    .severity {{ padding:2px 8px;border-radius:4px;font-size:10px;font-weight:bold }}
    .severity.error {{ background:#fee2e2;color:#b91c1c }}
    .severity.warning {{ background:#fef3c7;color:#92400e }}
    .severity.info {{ background:#dbeafe;color:#1e40af }}

    article h2 {{ font-size:14px;margin:0 0 6px;color:#0f172a }}
    .summary {{ color:#475569;margin:0 0 10px;font-size:11px }}

    .presentation {{ margin:8px 0 10px;padding:8px 10px;background:#f8fafc;border:1px solid #e2e8f0;border-radius:6px }}
    .presentation div {{ display:grid;grid-template-columns:78px 1fr;gap:8px;padding:3px 0 }}
    .presentation b {{ color:#475569;font-size:10px }} .presentation span {{ color:#1e293b;font-size:10px }}
    .supplement {{ margin:8px 0;padding:8px 10px;background:#eff6ff;border-radius:6px;font-size:10px }}
    .supplement p {{ margin:4px 0;color:#334155;white-space:pre-line }}

    blockquote {{ margin:6px 0;padding:8px 10px;background:#f8fafc;border-left:3px solid #cbd5e1;font-size:10px }}
    blockquote small {{ display:block;color:#64748b;margin-bottom:3px;font-weight:bold;font-size:10px }}

    .decision {{ margin-top:10px;padding:8px 10px;background:#f0fdf4;border-radius:6px;font-size:10px }}
    .decision.pending {{ background:#fffbeb }}
    .decision b {{ color:#166534 }} .decision.pending b {{ color:#92400e }}
    .decision p {{ margin:4px 0;color:#334155 }} .decision small {{ color:#94a3b8 }}
    </style><body>
    <h1>EMC 文档集审核报告</h1>
    <p class='meta'>任务：{html.escape(snapshot.run.set_id)} &nbsp;|&nbsp; 运行：{html.escape(snapshot.run.graph_id)} &nbsp;|&nbsp; 状态：{html.escape(snapshot.run.status)}</p>
    <div class='stats'>
    <div class='err'>需处理 {error_count} 项</div>
    <div class='warn'>需确认 {warning_count} 项</div>
    <div class='info'>提示 {info_count} 项</div>
    </div>
    {''.join(rows)}</body></html>"""


def generate_review_pdf(snapshot: GraphSnapshot) -> bytes:
    with tempfile.TemporaryDirectory(prefix="egraph-export-") as directory:
        html_path = os.path.join(directory, "review.html")
        pdf_path = os.path.join(directory, "review.pdf")
        with open(html_path, "w", encoding="utf-8") as stream:
            stream.write(_review_html(snapshot))
        args = [
            CHROME_PATH, "--headless", "--disable-gpu", "--no-sandbox",
            "--disable-software-rasterizer", f"--print-to-pdf={pdf_path}",
            f"file://{html_path}",
        ]
        subprocess.run(args, check=True, timeout=45, capture_output=True)
        with open(pdf_path, "rb") as stream:
            return stream.read()


def generate_evidence_package(snapshot: GraphSnapshot) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        # 1) Full machine-readable graph snapshot
        archive.writestr(
            "review-graph.json",
            json.dumps(_snapshot_dump_with_presentations(snapshot), ensure_ascii=False, indent=2),
        )
        # 2) Same user-friendly Excel as the standalone download
        archive.writestr("审核结论.xlsx", generate_review_excel(snapshot))

        # 3) Evidence catalog CSV (user-friendly Chinese headers)
        evidence_map = {item.evidence_id: item for item in snapshot.evidence}
        decisions_map = {item.finding_id: item for item in snapshot.decisions}
        findings_csv = io.StringIO()
        findings_writer = csv.writer(findings_csv)
        findings_writer.writerow([
            "序号", "问题类别", "严重程度", "问题标题", "问题说明",
            "涉及文档", "文件名", "页码/工作表", "原文摘录",
            "处理状态", "处理备注", "处理人", "处理时间",
            "系统状态", "核对对象", "已核对范围", "结构化发现", "影响", "LLM补充说明",
        ])
        for idx, finding in enumerate(snapshot.findings, 1):
            evidence_items = _readable_evidence_items(finding, evidence_map)
            decision = decisions_map.get(str(finding.finding_id))
            presentation = _finding_presentation(finding)
            for e in evidence_items:
                findings_writer.writerow([
                    idx,
                    CHECK_CATEGORY_LABELS.get(str(finding.check_id), str(finding.check_id)),
                    SEVERITY_LABELS.get(str(finding.severity), str(finding.severity)),
                    _plain_title(finding, evidence_items),
                    _plain_summary(finding, evidence_items),
                    DOC_TYPE_LABELS.get(getattr(e, "doc_type", ""), e.doc_type),
                    _display_filename(e),
                    _format_location(e),
                    _trim(getattr(e, "exact_quote", "")),
                    DECISION_LABELS.get(str(getattr(decision, "decision", "")), "待处理") if decision else "待处理",
                    _trim(getattr(decision, "comment", "")) if decision else "",
                    _trim(getattr(decision, "actor_id", "")) if decision else "",
                    _trim(getattr(decision, "created_at", "")) if decision else "",
                    presentation.get("status_label", ""),
                    presentation.get("subject", ""),
                    presentation.get("checked", ""),
                    presentation.get("issue", ""),
                    presentation.get("impact", ""),
                    _llm_supplement_text(presentation),
                ])
        archive.writestr("审核明细.csv", findings_csv.getvalue().encode("utf-8-sig"))
    return output.getvalue()
